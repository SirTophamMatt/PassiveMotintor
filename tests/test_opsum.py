"""Operational Summary: field map ↔ template agreement, rendering, storage,
carry-forward, image handling, the Intel gate on the image route, and the
scrub of the committed template (the repo may be public)."""
import io
import re
import zipfile

import pytest
from PIL import Image as PILImage
from pptx import Presentation
from pptx.oxml.ns import qn

from app import opsum, opsum_pptx

DATE = "2026-09-18"
TOKEN = re.compile(r"\{\{([a-z0-9_]+(?:\.\d+)*)\}\}")


def _png(w=400, h=200, colour=(200, 30, 30)):
    buf = io.BytesIO()
    PILImage.new("RGB", (w, h), colour).save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def images(monkeypatch, tmp_path):
    monkeypatch.setattr(opsum, "IMAGE_DIR", str(tmp_path / "opsum_images"))
    return tmp_path


def _all_text(xml_parts):
    return "".join(TOKEN.sub(lambda m: "{{%s}}" % m.group(1), x) for x in xml_parts)


def _template_tokens():
    z = zipfile.ZipFile(opsum.TEMPLATE_PATH)
    toks = set()
    for n in z.namelist():
        if n.endswith(".xml") and (n.startswith("ppt/slides/") or n.startswith("ppt/slideLayouts/")):
            # Tokens can only be read from joined run text; the template keeps
            # each token inside one run, so the raw XML is enough.
            toks |= set(TOKEN.findall(z.read(n).decode("utf8")))
    return toks


def _shape_names():
    prs = Presentation(opsum.TEMPLATE_PATH)
    names = set()
    for s in prs.slides:
        for el in s._element.iter(qn("p:cNvPr")):
            names.add(el.get("name"))
    return names, prs


def _render(data=None, **kw):
    data = data or opsum.blank(DATE)
    return Presentation(io.BytesIO(opsum_pptx.render(DATE, data, **kw)))


def _xml(prs_bytes):
    z = zipfile.ZipFile(io.BytesIO(prs_bytes))
    return {n: z.read(n) for n in z.namelist()}


# --------------------------------------------------------------------------- #
# The field map and the template cannot drift apart
# --------------------------------------------------------------------------- #
def test_every_field_has_a_place_in_the_template():
    toks = _template_tokens()
    names, _ = _shape_names()
    for key, f in opsum.FIELDS.items():
        if f.kind in ("text", "line"):
            assert key in toks, key
        elif f.kind in ("table", "colors"):
            for i in range(len(f.rows)):
                for j in range(len(f.cols)):
                    assert f"{key}.{i}.{j}" in toks, (key, i, j)
        elif f.kind == "grid":
            assert f"OPS_GRID:{key}" in names, key
        elif f.kind == "image":
            assert f"OPS_IMG:{key}" in names, key


def test_every_template_token_is_known():
    known = set(opsum.derived_tokens(DATE))
    for key, f in opsum.FIELDS.items():
        known.add(key)
        for i in range(len(f.rows)):
            for j in range(len(f.cols)):
                known.add(f"{key}.{i}.{j}")
    assert _template_tokens() - known == set()


def test_grid_labels_match_the_template_cells():
    _, prs = _shape_names()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name.startswith("OPS_GRID:"):
                key = shape.name.split(":", 1)[1]
                cells = {opsum_pptx._norm(c.text) for r in shape.table.rows
                         for c in r.cells if c.text.strip()}
                assert cells == {opsum_pptx._norm(x) for x in opsum.FIELDS[key].labels}


def test_optional_slides_exist_in_template():
    _, prs = _shape_names()
    marked = {s._element.cSld.get("name").split(":", 1)[1] for s in prs.slides
              if (s._element.cSld.get("name") or "").startswith("OPS_OPTIONAL:")}
    assert marked == set(opsum.OPTIONAL_SLIDES)


# --------------------------------------------------------------------------- #
# The committed template holds no personal or day content
# --------------------------------------------------------------------------- #
def test_template_is_scrubbed():
    z = zipfile.ZipFile(opsum.TEMPLATE_PATH)
    names = z.namelist()
    assert not [n for n in names if re.search(r"comment|changesInfo|revisionInfo|customXml|custom\.xml", n)]
    blob = b"".join(z.read(n) for n in names if not n.startswith("ppt/media/"))
    for needle in (b"isentia", b"userId=", b"DJCS"):
        assert needle not in blob
    prs = Presentation(opsum.TEMPLATE_PATH)
    assert prs.core_properties.author == "" and prs.core_properties.last_modified_by == ""
    for s in prs.slides:
        if s.has_notes_slide:
            assert not s.notes_slide.notes_text_frame.text.strip()


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #
def test_empty_draft_renders_clean(images):
    out = opsum_pptx.render(DATE, opsum.blank(DATE))
    parts = _xml(out)
    assert zipfile.ZipFile(io.BytesIO(out)).testzip() is None
    leftovers = [n for n, b in parts.items() if n.endswith(".xml") and b"{{" in b]
    assert leftovers == []
    prs = Presentation(io.BytesIO(out))
    assert len(prs.slides) == 12 - len(opsum.OPTIONAL_SLIDES)
    # No image slot left as a placeholder, no slide link to a removed slide.
    for s in prs.slides:
        assert not [el for el in s._element.iter(qn("p:cNvPr"))
                    if el.get("name", "").startswith("OPS_IMG:")]
        for rel in s.part.rels.values():
            if rel.reltype.endswith("/slide"):
                assert rel.target_part in {x.part for x in prs.slides}


def test_optional_slides_included_on_request(images):
    data = opsum.blank(DATE)
    for k in opsum.OPTIONAL_SLIDES:
        data["optional"][k] = True
    assert len(_render(data).slides) == 12


def test_menu_tile_jumps_to_activation_slide():
    prs = Presentation(opsum.TEMPLATE_PATH)
    activation = prs.slides[8].part
    targets = {rel.target_part for rel in prs.slides[0].part.rels.values()
               if rel.reltype.endswith("/slide")}
    assert activation in targets


def test_header_and_inline_values(images):
    data = opsum.blank(DATE)
    data["fields"]["scc_class1"] = "Monitoring"
    data["fields"]["warnings"] = [["1", "2", "3", "4"]]
    prs = _render(data)
    lay = [l for l in prs.slide_layouts if l.name == "Title Slide"][0]
    texts = [sh.text_frame.text for sh in lay.shapes if sh.has_text_frame]
    assert "Current as at 0830hrs on Friday 18 September 2026" in texts
    assert any("Class 1: Monitoring" in t for t in texts)
    counts = sorted((sh for sh in prs.slides[1].shapes
                     if sh.name == "Rectangle: Rounded Corners 71"), key=lambda s: s.left)
    assert [c.text_frame.text for c in counts] == ["1", "2", "3", "4"]


def test_derived_dates():
    tok = opsum.derived_tokens(DATE)
    assert tok["av_days.0"] == "12/09" and tok["av_days.6"] == "18/09"
    assert tok["code1_days.0"] == "10/09" and tok["code1_days.6"] == "16/09"
    assert (tok["year_this"], tok["year_last"]) == ("2026", "2025")


def test_text_field_expansion(images):
    data = opsum.blank(DATE)
    data["fields"]["kp_last24"] = ("First **bold** point\n  sub point\n\n"
                                   "# A heading\n- dash is stripped\n"
                                   "see [the site](https://example.org/x)\n\n\n")
    out = opsum_pptx.render(DATE, data)
    prs = Presentation(io.BytesIO(out))
    tf = [sh for sh in prs.slides[1].shapes if sh.name == "Table 3"][0].table.cell(1, 0).text_frame
    texts = [p.text for p in tf.paragraphs]
    i = texts.index("First bold point")
    assert texts[i:i + 6] == ["First bold point", "sub point", "", "A heading",
                              "dash is stripped", "see the site"]
    paras = tf.paragraphs
    first, sub, spacer, heading = paras[i], paras[i + 1], paras[i + 2], paras[i + 3]
    assert [r.font.bold for r in first.runs] == [False, True, False]
    assert sub._p.pPr.get("lvl") == "1"
    assert int(sub._p.pPr.get("marL")) > int(first._p.pPr.get("marL"))
    assert first._p.pPr.find(qn("a:buChar")) is not None
    assert heading._p.pPr.find(qn("a:buChar")) is None and heading.runs[0].font.bold
    # an empty spacer keeps the body size instead of the 18 pt default
    assert spacer._p.find(qn("a:endParaRPr")).get("sz") == "700"
    link = paras[i + 5].runs[-1]
    assert link.hyperlink.address == "https://example.org/x"
    # trailing blank lines add nothing
    data["fields"]["kp_last24"] = data["fields"]["kp_last24"].rstrip("\n")
    again = Presentation(io.BytesIO(opsum_pptx.render(DATE, data)))
    tf2 = [sh for sh in again.slides[1].shapes if sh.name == "Table 3"][0].table.cell(1, 0).text_frame
    assert len(tf2.paragraphs) == len(tf.paragraphs)


def test_value_without_markup_keeps_template_weight(images):
    data = opsum.blank(DATE)
    data["fields"]["warnings"] = [["7", "", "", ""]]
    prs = _render(data)
    count = min((sh for sh in prs.slides[1].shapes
                 if sh.name == "Rectangle: Rounded Corners 71"), key=lambda s: s.left)
    assert count.text_frame.paragraphs[0].runs[0].font.bold is True


def test_bold_inline_value_in_table_cell(images):
    data = opsum.blank(DATE)
    data["optional"]["planned_burning"] = True
    data["fields"]["pb_ytd"][2][0] = "**35,598** (18.2% of program)*"
    prs = _render(data)
    slide = [s for s in prs.slides
             if (s._element.cSld.get("name") or "") == "OPS_OPTIONAL:planned_burning"][0]
    cell = [sh for sh in slide.shapes if sh.name == "Table 2"][0].table.cell(3, 1)
    runs = cell.text_frame.paragraphs[0].runs
    assert [(r.text, bool(r.font.bold)) for r in runs] == [
        ("35,598", True), (" (18.2% of program)*", False)]


def _slide_with(prs, shape_name):
    return [s for s in prs.slides if any(sh.name == shape_name for sh in s.shapes)][0]


def _fill(cell):
    f = cell._tc.tcPr.find(qn("a:solidFill")) if cell._tc.tcPr is not None else None
    return None if f is None else f[0].get("val")


def test_activation_grid_uses_legend_fills(images):
    data = opsum.blank(DATE)
    data["fields"]["rcc"]["RCC – EMR"] = "Active"
    data["fields"]["icc"]["Seymour"] = "Readiness"
    prs = _render(data)
    slide = _slide_with(prs, "OPS_GRID_LEGEND")
    legend = {c.text.strip(): c for r in [s for s in slide.shapes
              if s.name == "OPS_GRID_LEGEND"][0].table.rows for c in r.cells}
    cells = {c.text.strip(): c for name in ("OPS_GRID:rcc", "OPS_GRID:icc")
             for r in [s for s in slide.shapes if s.name == name][0].table.rows
             for c in r.cells}
    assert _fill(cells["RCC – EMR"]) == _fill(legend["Active"]) == "00B050"
    seymour = cells["Seymour"]._tc.tcPr.find(qn("a:solidFill"))
    readiness = legend["Readiness"]._tc.tcPr.find(qn("a:solidFill"))
    assert seymour[0].tag == readiness[0].tag and seymour[0].get("val") == readiness[0].get("val")
    assert cells["Ararat"]._tc.tcPr.find(qn("a:solidFill")) is None      # Not Active


def test_workload_colours_and_flood_status(images):
    data = opsum.blank(DATE)
    data["fields"]["av_workload"][0][3] = "Orange"
    data["optional"]["flood"] = True
    data["fields"]["flood_gauges"][2] = ["minor", "MODERATE", "", "MAJOR", "other"]
    prs = _render(data)
    wl = [s for s in prs.slides[4].shapes if s.name == "Table 4"][0].table
    assert _fill(wl.cell(2, 4)) == "ED7D31"
    assert _fill(wl.cell(2, 1)) == "E9EBF5"          # untouched = no escalation
    assert wl.cell(2, 4).text == ""
    flood = [s for s in prs.slides
             if (s._element.cSld.get("name") or "") == "OPS_OPTIONAL:flood"][0]
    g = [s for s in flood.shapes if s.name == "Table 11"][0].table
    assert [_fill(g.cell(3, j)) for j in range(1, 6)] == [
        "00B050", "ED7D31", "F2F2F2", "FF0000", "F2F2F2"]


def test_image_fits_inside_its_frame(images):
    tpl = Presentation(opsum.TEMPLATE_PATH)
    frame = [s for s in _slide_with(tpl, "OPS_IMG:scc_roster").shapes
             if s.name == "OPS_IMG:scc_roster"][0]
    data = opsum.blank(DATE)
    data["fields"]["scc_roster"] = opsum.store_image(_png(100, 400), "tall.png")
    out = _render(data)
    pic = [s for s in _slide_with(out, "OPS_IMG:scc_roster").shapes
           if s.name == "OPS_IMG:scc_roster"][0]
    assert pic.height == pytest.approx(frame.height, abs=2)
    assert pic.width == pytest.approx(frame.height / 4, rel=0.01)       # aspect kept
    assert frame.left <= pic.left and pic.left + pic.width <= frame.left + frame.width
    assert pic.image.blob == _png(100, 400)


# --------------------------------------------------------------------------- #
# Model + storage
# --------------------------------------------------------------------------- #
def test_normalise_forces_shape_and_values():
    data = opsum.normalise({
        "fields": {
            "warnings": [["1", "2"], ["x"]],
            "av_workload": [["Orange", "Purple"]],
            "rcc": {"RCC – BSW": "Active", "RCC – GMP": "On fire", "Bogus": "Active"},
            "kp_today": float("nan"),
            "as_at": "a\nb",
            "unknown": "dropped",
            "readiness": {"sha": "../../etc/passwd", "ext": "png"},
        },
        "optional": {"flood": 1, "nonsense": True},
    }, DATE)
    f = data["fields"]
    assert f["warnings"] == [["1", "2", "", ""]]
    assert f["av_workload"][0][:2] == ["Orange", ""]
    assert f["rcc"]["RCC – BSW"] == "Active" and f["rcc"]["RCC – GMP"] == "Not Active"
    assert "Bogus" not in f["rcc"]
    assert f["kp_today"] == "" and f["as_at"] == "a b"
    assert "unknown" not in f and f["readiness"] is None
    assert data["optional"] == {"planned_burning": False, "avian": False, "flood": True}


def test_save_load_versions_and_issue(db):
    data = opsum.blank(DATE)
    data["fields"]["kp_today"] = "Fine"
    assert opsum.save(DATE, data) == 1
    assert opsum.save(DATE, data) == 2
    loaded, meta = opsum.load(DATE)
    assert loaded["fields"]["kp_today"] == "Fine" and meta["status"] == "draft"
    opsum.issue(DATE, data)
    _, meta = opsum.load(DATE)
    assert meta["status"] == "issued" and meta["issued_at"]
    from app import database
    assert len(database.read_df("SELECT * FROM opsum_versions")) == 1


def test_carry_forward_only_copies_carry_fields(db):
    prev = opsum.blank("2026-09-17")
    prev["fields"]["src_current"] = "Controller A (SES)"
    prev["fields"]["kp_last24"] = "Yesterday's news"
    prev["fields"]["rcc"]["RCC – GIP"] = "Active"
    prev["optional"]["avian"] = True
    opsum.save("2026-09-17", prev)
    opsum.save("2026-09-10", opsum.blank("2026-09-10"))      # older, ignored
    data, meta, carried_from = opsum.open_draft(DATE)
    assert meta is None and str(carried_from) == "2026-09-17"
    assert data["fields"]["src_current"] == "Controller A (SES)"
    assert data["fields"]["rcc"]["RCC – GIP"] == "Active"
    assert data["fields"]["kp_last24"] == ""
    assert data["optional"]["avian"] is True
    assert data["fields"]["as_at"] == "0830hrs on Friday 18 September 2026"


def test_image_store_validates_by_content(images):
    ref = opsum.store_image(_png(), "../../evil name.png")
    assert ref["ext"] == "png" and ref["name"] == "evil name.png"
    assert opsum.image_path(ref)
    assert opsum.store_image(_png(), "again.png")["sha"] == ref["sha"]      # dedupe
    with pytest.raises(ValueError):
        opsum.store_image(b"<svg onload=alert(1)>", "x.png")
    with pytest.raises(ValueError):
        opsum.store_image(b"", "x.png")
    assert opsum.image_file(f"{ref['sha']}.png")
    for bad in ("../x.png", f"{ref['sha']}.svg", f"{ref['sha']}.png/../x", "x"):
        assert opsum.image_file(bad) is None


# --------------------------------------------------------------------------- #
# Page wiring
# --------------------------------------------------------------------------- #
def test_collect_rebuilds_a_draft():
    from app.pages import opsum as page
    data = page.collect(
        [{"key": "kp_today"}], ["Sunny"],
        [{"key": "warnings"}], [[{"label": "Count", "c0": "1", "c1": "2", "c2": "", "c3": "4"}]],
        [{"key": "rcc"}], [[{"label": "RCC – BSW", "state": "Active"}]],
        [{"key": "readiness"}], [None],
        [{"key": "flood"}], [["on"]])
    norm = opsum.normalise(data, DATE)
    assert norm["fields"]["kp_today"] == "Sunny"
    assert norm["fields"]["warnings"] == [["1", "2", "", "4"]]
    assert norm["fields"]["rcc"]["RCC – BSW"] == "Active"
    assert norm["optional"]["flood"] is True


def test_default_password_keeps_summary_closed(monkeypatch):
    from app.pages import intel
    from app.pages import opsum as page
    monkeypatch.setattr(intel, "DESKTOP", False)
    monkeypatch.setattr(intel, "INTEL_PASSWORD", intel.DEFAULT_PASSWORD)
    assert not page.available()
    monkeypatch.setattr(intel, "INTEL_PASSWORD", "something-else")
    assert page.available()


def test_image_route_requires_the_intel_session(images, monkeypatch):
    import dash
    from app.pages import intel
    from app.pages import opsum as page
    monkeypatch.setattr(intel, "INTEL_PASSWORD", "something-else")
    ref = opsum.store_image(_png(), "a.png")
    app = dash.Dash(__name__)
    app.layout = dash.html.Div()
    app.server.secret_key = "test"
    page.register_callbacks(app)
    client = app.server.test_client()
    url = f"{page.IMAGE_ROUTE}/{ref['sha']}.png"
    assert client.get(url).status_code == 403
    with client.session_transaction() as sess:
        sess[intel.SESSION_KEY] = True
    r = client.get(url)
    assert r.status_code == 200 and r.data == _png()
    # A name that reaches the route but is not <sha1>.<ext> is refused ...
    assert client.get(f"{page.IMAGE_ROUTE}/..config.json").status_code == 404
    assert client.get(f"{page.IMAGE_ROUTE}/{ref['sha']}.svg").status_code == 404
    # ... and a path with a slash never reaches it (Dash's own page answers).
    r = client.get(f"{page.IMAGE_ROUTE}/..%2f..%2fconfig.json")
    assert r.mimetype == "text/html"


def test_blank_intel_password_env_means_default():
    from app.pages import intel
    assert intel.password_from_env(None) == intel.DEFAULT_PASSWORD
    assert intel.password_from_env("") == intel.DEFAULT_PASSWORD
    assert intel.password_from_env("   ") == intel.DEFAULT_PASSWORD
    assert intel.password_from_env("s3cret") == "s3cret"
