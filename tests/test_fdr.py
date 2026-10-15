"""Fire Danger Ratings tool (app/modules/fdr): parsing, history, resolution,
and the workbook writer.

Feed fixture ``cfa_tfbfdr_sample.xml`` is CONSTRUCTED test data in the CFA
feed's structure, not a real forecast. The workbook is built here to the
layout described for the FDR Daily Statistics workbook (TFB C, Multi Day D,
Date E, districts G–O, Comments P, data from row 12, 1 Jul 2026 – 30 Jun
2027), plus an Excel-extension dropdown injected into the XML — the kind of
thing an openpyxl save would silently drop.
"""
import os
import re
import zipfile
from datetime import date, datetime, timedelta

import openpyxl
import pytest
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment

from app import database
from app.modules.fdr import model, parse, resolve, scraper, store, workbook
from tests.conftest import load_fixture

FEED = load_fixture("cfa_tfbfdr_sample.xml")
D15, D16, D17, D18 = (date(2026, 10, d) for d in (15, 16, 17, 18))


def at(d, hh, mm=0):
    return datetime(d.year, d.month, d.day, hh, mm, tzinfo=model.TZ)


def fetcher(body):
    return lambda url, timeout: body


def run(body=FEED, when=None):
    return scraper.run("test", fetcher=fetcher(body), when=when or at(D15, 5, 30))


# --------------------------------------------------------------------------- #
# Model
# --------------------------------------------------------------------------- #
def test_districts_in_never_double_counts_nested_names():
    assert model.districts_in("North Central and Central") == ["North Central", "Central"]
    assert model.districts_in("West and South Gippsland, South West") == [
        "West and South Gippsland", "South West"]
    assert model.districts_in("East Gippsland") == ["East Gippsland"]


@pytest.mark.parametrize("label,expected", [
    ("W&S Gippsland", "West and South Gippsland"),
    ("West & South Gippsland", "West and South Gippsland"),
    ("Nth Central", "North Central"), ("NORTHERN COUNTRY", "Northern Country"),
    ("South West", "South West"), ("Central Comments", None), ("Date", None)])
def test_header_aliases(label, expected):
    assert model.normalise_district(label) == expected


@pytest.mark.parametrize("text,ratings,ok", [
    ("HIGH", ["High"], True),
    ("No Rating", ["No Rating"], True),
    ("High and Extreme", ["Extreme", "High"], True),
    ("HIGH/EXTREME", ["Extreme", "High"], True),
    ("VERY HIGH", [], False),            # legacy word: never mapped to High
    ("SEVERE", [], False),
    ("HIGH - see cfa.vic.gov.au", ["High"], True),
    ("HIGH - except the north", ["High"], False),   # qualifies where: review
    ("", [], False), ("TBA", [], False)])
def test_parse_ratings(text, ratings, ok):
    assert model.parse_ratings(text) == (ratings, ok)


# --------------------------------------------------------------------------- #
# Parser
# --------------------------------------------------------------------------- #
def test_parse_normal_day():
    res = parse.parse_feed(FEED, fetched_at=at(D15, 5, 30))
    day = res.days[0]
    assert day.valid_date == D15
    assert len(day.ratings) == 9
    assert day.ratings["Mallee"].ratings == ["Extreme"]
    assert day.ratings["West and South Gippsland"].ratings == ["Moderate"]
    assert day.tfb["Mallee"].status == "YES" and day.tfb["Central"].status == "NO"
    assert day.issued.isoformat() == "2026-10-15T05:30:00+11:00"   # BoM issue time


def test_statewide_no_tfb_statement_is_explicit_no():
    d16 = parse.parse_feed(FEED, fetched_at=at(D15, 5, 30)).days[1]
    assert {t.status for t in d16.tfb.values()} == {"NO"}
    assert len(d16.tfb) == 9
    assert d16.tfb["Mallee"].basis == "statewide statement"


def test_missing_tfb_section_is_unknown_not_no():
    d17 = parse.parse_feed(FEED, fetched_at=at(D15, 5, 30)).days[2]
    assert d17.tfb == {}
    assert "South West" not in d17.ratings       # missing district stays missing


def test_multiple_ratings_and_unrecognised_values():
    days = parse.parse_feed(FEED, fetched_at=at(D15, 5, 30)).days
    assert days[1].ratings["Central"].ratings == ["Extreme", "High"]
    nc = days[3].ratings["North Central"]
    assert nc.ratings == [] and not nc.recognised


def test_weekday_mismatch_rejects_the_day():
    bad = FEED.replace(b"Friday, 16 October 2026", b"Saturday, 16 October 2026")
    res = parse.parse_feed(bad, fetched_at=at(D15, 5, 30))
    assert D16 not in [d.valid_date for d in res.days]
    assert any("does not fall on" in p for p in res.problems)


def test_implausible_dates_are_ignored():
    res = parse.parse_feed(FEED, fetched_at=at(date(2027, 1, 1), 9))
    assert res.days == [] and res.problems


def test_html_variants_and_district_hint():
    lines = parse.html_lines("<table><tr><td>Mallee</td><td>HIGH</td></tr></table>"
                             "<div><strong>Central:</strong> EXTREME<br>North East: NO RATING</div>")
    assert lines == ["Mallee: HIGH", "Central: EXTREME", "North East: NO RATING"]
    ratings, tfb, _ = parse.parse_day_text(
        ["Total Fire Ban: YES", "Fire Danger Rating: EXTREME"], district_hint="Mallee")
    assert ratings["Mallee"].ratings == ["Extreme"]
    assert tfb["Mallee"].status == "YES"


def test_garbage_feed_reports_a_problem():
    res = parse.parse_feed(b"<html><body>Service unavailable</body></html>")
    assert res.days == []


LIVE = load_fixture("cfa_tfbfdr_live_20261015.xml")
LIVE_CENTRAL = load_fixture("cfa_central_live_20261015.xml")


def test_live_feed_captured_15_oct_2026():
    """The real CFA statewide feed: 5 items, the 5th a placeholder."""
    res = parse.parse_feed(LIVE, fetched_at=at(D15, 15))
    assert [d.valid_date for d in res.days] == [D15 + timedelta(days=i) for i in range(5)]
    assert [d.unissued for d in res.days] == [False] * 4 + [True]
    assert res.days[0].issued.isoformat() == "2026-10-15T05:30:00+11:00"
    today = res.days[0]
    assert {d: r.ratings[0] for d, r in today.ratings.items()} == {
        "Central": "Moderate", "East Gippsland": "Moderate", "Mallee": "Moderate",
        "North Central": "Moderate", "North East": "Moderate",
        "Northern Country": "No Rating", "South West": "No Rating",
        "West and South Gippsland": "Moderate", "Wimmera": "No Rating"}
    assert all(r.recognised for r in today.ratings.values())
    assert {t.status for t in today.tfb.values()} == {"NO"} and len(today.tfb) == 9
    assert today.tfb["Mallee"].raw == "Mallee: NO - RESTRICTIONS MAY APPLY"


def test_live_placeholder_day_is_not_stored(db):
    s = scraper.run("test", fetcher=fetcher(LIVE), when=at(D15, 15))
    assert s["status"] == "ok" and s["inserted"] == 36          # 4 issued days x 9
    assert store.forecasts(date(2026, 10, 19), date(2026, 10, 19)).empty
    assert store.open_flags().empty


def test_live_district_feed_legend_is_not_a_ban():
    """Every district item ends 'Displays when Total Fire Ban in force' (an
    icon legend) — it must never read as a declaration."""
    res = parse.parse_feed(LIVE_CENTRAL, district_hint="Central", fetched_at=at(D15, 15))
    day = res.days[0]
    assert day.ratings["Central"].ratings == ["Moderate"]
    assert day.tfb["Central"].status == "NO"
    assert list(day.ratings) == ["Central"]


# --------------------------------------------------------------------------- #
# Collection + history
# --------------------------------------------------------------------------- #
def test_first_run_stores_every_value(db):
    s = run()
    assert s["status"] == "ok"
    # 9 + 9 + 8 + 9 district values
    assert s["inserted"] == 35 and s["retrieved"] == 35
    df = store.forecasts(D15, D18, current_only=True)
    assert set(df["district"]) == set(model.DISTRICTS)
    assert df[(df.valid_date == "2026-10-16") & (df.district == "Central")
              ]["multi"].iloc[0] == 1


def test_duplicate_retrieval_adds_nothing(db):
    run()
    before = len(store.forecasts(D15, D18))
    s = run(when=at(D15, 6, 0))
    assert s["inserted"] == 0 and s["updated"] == 0 and s["unchanged"] == 35
    assert len(store.forecasts(D15, D18)) == before
    row = store.forecasts(D15, D15, current_only=True).iloc[0]
    assert row["seen_count"] == 2 and row["last_seen_at"].startswith("2026-10-15T06:00")


def test_amendment_is_a_new_revision(db):
    run()
    amended = FEED.replace(b"Central: HIGH&lt;br&gt;", b"Central: EXTREME&lt;br&gt;", 1)
    s = run(amended, when=at(D15, 13, 40))
    assert s["updated"] == 1
    hist = store.forecasts(D15, D15)
    central = hist[hist.district == "Central"].sort_values("revision")
    assert list(central["rating"]) == ["High", "Extreme"]
    assert list(central["is_current"]) == [0, 1]


def test_unrecognised_value_never_replaces_a_good_one(db):
    run()
    garbled = FEED.replace(b"Mallee: EXTREME&lt;br&gt;", b"Mallee: TBA&lt;br&gt;", 1)
    run(garbled, when=at(D15, 7))
    cur = store.forecasts(D15, D15, current_only=True)
    assert cur[cur.district == "Mallee"]["rating"].iloc[0] == "Extreme"
    assert (store.open_flags()["kind"] == "unrecognised").any()


def test_source_outage_never_erases_data(db):
    run()
    before = store.forecasts(D15, D18).fillna("").to_dict("records")

    def down(url, timeout):
        raise ConnectionError("CFA unreachable")
    s = scraper.run("test", fetcher=down, when=at(D15, 8))
    assert s["status"] == "failed" and "CFA unreachable" in s["errors"][0]
    assert store.forecasts(D15, D18).fillna("").to_dict("records") == before
    assert store.last_run()["status"] == "failed"


def test_unreadable_feed_saves_raw_body_and_fails(db, tmp_path, monkeypatch):
    monkeypatch.setattr(scraper, "DEBUG_DIR", str(tmp_path))
    s = run(b"<rss><channel><item><title>x</title></item></channel></rss>")
    assert s["status"] == "failed"
    assert (tmp_path / "fdr_debug_cfa-state.xml").exists()
    assert store.forecasts(D15, D18).empty


def test_flags_are_not_repeated_by_later_runs(db):
    store.add_flag(None, None, None, "feed-format", "odd item")
    store.add_flag(None, None, None, "feed-format", "odd item")
    run()
    n = len(store.open_flags())
    run(when=at(D15, 6))
    assert len(store.open_flags()) == n
    assert (store.open_flags()["detail"] == "odd item").sum() == 1


def test_missing_district_is_flagged(db):
    run()
    flags = store.open_flags()
    assert ((flags.kind == "missing") & (flags.district == "South West")).any()


# --------------------------------------------------------------------------- #
# Resolution
# --------------------------------------------------------------------------- #
def test_confirmed_day_record(db):
    run()
    rec = resolve.resolve_day(D15, now=at(D15, 9))
    assert rec.complete
    assert rec.multi == "NO"
    assert rec.tfb_any == "YES" and rec.tfb_districts == ["Mallee", "Wimmera"]


def test_future_forecast_is_preliminary_until_seen_near_the_day(db):
    run()   # read on the 15th at 05:30
    rec = resolve.resolve_day(D17, now=at(D17, 9))
    assert rec.districts["Mallee"].state == "preliminary"
    assert rec.multi is None                      # never assumed NO
    run(when=at(D16, 20))                         # re-read the evening before
    rec = resolve.resolve_day(D17, now=at(D17, 9))
    assert rec.districts["Mallee"].state == "confirmed"
    assert rec.districts["South West"].state == "none"
    assert rec.multi is None                      # one district still unknown
    assert rec.tfb_any is None                    # no TFB published: not NO


def test_multi_from_published_values(db):
    run(when=at(D15, 20))
    rec = resolve.resolve_day(D16, now=at(D16, 9))
    assert rec.multi == "YES"
    assert rec.districts["Central"].multi_detail == "High and Extreme (published)"
    assert rec.tfb_any == "NO"


def test_multi_from_an_amendment_during_the_day(db):
    run(when=at(D15, 5))
    # CFA re-issues: the day's "forecast issued at" moves to 1:30 PM. The
    # tool sees it at 13:40, but the change took effect when BoM issued it.
    amended = FEED.replace(b"Central: HIGH&lt;br&gt;", b"Central: EXTREME&lt;br&gt;", 1) \
        .replace(b"15 October 2026 05:30 AM", b"15 October 2026 01:30 PM", 1)
    run(amended, when=at(D15, 13, 40))
    rec = resolve.resolve_day(D15, now=at(D15, 18))
    c = rec.districts["Central"]
    assert c.multi and rec.multi == "YES"
    assert c.ratings == ["Extreme", "High"]
    assert c.multi_detail == "High until 13:30, Extreme from 13:30 (amended on the day)"


def test_overnight_revision_finalises_rather_than_multi(db):
    # Forecast read the evening before, revised at 04:00 on the day.
    early = FEED.replace(b"Central: HIGH&lt;br&gt;", b"Central: MODERATE&lt;br&gt;", 1)
    run(early, when=at(D15 - timedelta(days=1), 18))
    run(when=at(D15, 4))
    rec = resolve.resolve_day(D15, now=at(D15, 9))
    assert not rec.districts["Central"].multi
    assert rec.districts["Central"].ratings == ["High"]


def test_conflicting_sources_are_flagged_not_chosen(db):
    run()
    district_body = (b"<rss><channel><item><title>Thursday, 15 October 2026</title>"
                     b"<description><![CDATA[<p>Fire Danger Ratings</p>"
                     b"<p>Central: EXTREME</p>]]></description></item></channel></rss>")
    conn = store.connect()
    res = parse.parse_feed(district_body, district_hint="Central",
                           fetched_at=at(D15, 6))
    scraper._store_feed(conn, scraper.DISTRICT, "x", "Central", district_body, res,
                        at(D15, 6), None, {"retrieved": 0, "inserted": 0,
                                           "updated": 0, "unchanged": 0})
    conn.commit()
    conn.close()
    rec = resolve.resolve_day(D15, now=at(D15, 9))
    assert rec.districts["Central"].state == "conflict"
    assert any("Central" in r for r in rec.review)
    resolve.flag_conflicts(at(D15, 9))
    assert (store.open_flags()["kind"] == "conflict").any()


def test_overrides_sit_on_top_and_can_be_revoked(db):
    run()
    ok, _ = store.add_override(D15, "Central", "rating", "High/Extreme", "CFA amendment")
    assert ok
    rec = resolve.resolve_day(D15, now=at(D15, 9))
    assert rec.districts["Central"].source == "override" and rec.multi == "YES"
    assert store.forecasts(D15, D15, current_only=True).query(
        "district == 'Central'")["rating"].iloc[0] == "High"   # source untouched
    oid = int(store.overrides().iloc[0]["id"])
    store.revoke_override(oid)
    assert resolve.resolve_day(D15, now=at(D15, 9)).multi == "NO"


@pytest.mark.parametrize("args,err", [
    (("2026-10-15", "Central", "rating", "Very High", "x"), "not an AFDRS"),
    (("2026-10-15", "Central", "rating", "High", ""), "reason"),
    (("2026-10-15", "Nowhere", "tfb", "YES", "x"), "nine districts"),
    (("2026-10-15", "Central", "multi", "YES", "x"), "whole day"),
    (("15/10/2026x", "*", "multi", "YES", "x"), "not a date")])
def test_override_validation(db, args, err):
    ok, msg = store.add_override(*args)
    assert not ok and err in msg


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #
# The real workbook's validations: all three dropdowns are Excel extensions
# (lists on the "List" sheet), exactly as in FDR Daily Statistics 2026-27.
X14_EXT = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations count="3" '
    'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" '
    'showErrorMessage="1" promptTitle="Rating" prompt="Input rating"><x14:formula1>'
    '<xm:f>List!$A$3:$A$%d</xm:f></x14:formula1>'
    '<xm:sqref>G109:P109 G12:O108 G110:O376</xm:sqref></x14:dataValidation>'
    '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" '
    'showErrorMessage="1"><x14:formula1><xm:f>List!$B$3:$B$5</xm:f></x14:formula1>'
    '<xm:sqref>C12:C376</xm:sqref></x14:dataValidation>'
    '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" '
    'showErrorMessage="1"><x14:formula1><xm:f>List!$C$3:$C$5</xm:f></x14:formula1>'
    '<xm:sqref>D12:D376</xm:sqref></x14:dataValidation></x14:dataValidations>'
    '</ext></extLst>')
# Column order of the real workbook (NOT CFA's order).
WB_DISTRICTS = ["Mallee", "Wimmera", "South West", "Northern Country", "North Central",
                "Central", "North East", "West and South Gippsland", "East Gippsland "]
COL = dict(zip([d.strip() for d in WB_DISTRICTS], "GHIJKLMNO"))
PICK = "Pick Rating\u2026."


def build_workbook(path, *, ratings=("No Rating", "Moderate", "High", "Extreme",
                                     "Catastrophic")):
    """A workbook laid out like the real FDR Daily Statistics 2026-27: a
    summary table (rows 2-8) that ALSO names every district, the data header at
    row 11, real district order, YES/NO, a "Pick Rating…." prompt pre-filled in
    days not yet entered, a cell note, and extension dropdowns. Values are
    invented."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["C2"], ws["E2"], ws["F2"] = "Rating", "Totals", "%"
    for i, name in enumerate(WB_DISTRICTS):
        ws.cell(row=2, column=7 + i, value=name)
    for k, rating in enumerate(model.RATINGS):
        r = 3 + k
        ws[f"C{r}"], ws[f"D{r}"] = rating, k + 1
        ws[f"E{r}"] = f"=SUM(G{r}:O{r})"
        for col in "GHIJKLMNO":
            ws[f"{col}{r}"] = f"=COUNTIF({col}12:INDEX({col}:{col},ROWS({col}:{col})), C{r})"
    heads = {"C": "TFB declared (Yes/No)", "D": "Multi Day", "E": "Date", "F": "Day",
             "P": "Comments"}
    for col, h in heads.items():
        ws[f"{col}11"] = h
    ws["D11"].comment = Comment("YES/NO - did the FDR change during the day", "SCC")
    for i, name in enumerate(WB_DISTRICTS):
        ws.cell(row=11, column=7 + i, value=name)
    d = date(2026, 7, 1)
    r = 12
    while d <= date(2027, 6, 30):
        # First date typed; the rest are =E(n-1)+1 (no cached value — openpyxl
        # does not calculate), which the reader must follow.
        ws[f"E{r}"] = datetime(d.year, d.month, d.day) if r == 12 else f"=E{r - 1}+1"
        ws[f"E{r}"].number_format = "dd/mm/yyyy"
        ws[f"F{r}"] = f'=TEXT(E{r},"ddd")'
        if d >= D15:                       # not yet entered: the prompt value
            for col in "GHIJKLMNO":
                ws[f"{col}{r}"] = PICK
        d += timedelta(days=1)
        r += 1
    lists = wb.create_sheet("List")
    lists["A2"], lists["B2"], lists["C2"] = ("Rating", "TFB declared (YES/NO)",
                                             "Multi Day (YES/NO)")
    for i, v in enumerate((PICK,) + tuple(ratings), 3):
        lists[f"A{i}"] = v
    lists["B3"], lists["B4"] = "YES", "NO"
    lists["C3"], lists["C4"] = "YES", "NO"
    ws.conditional_formatting.add(
        "G12:O376", CellIsRule(operator="equal", formula=['"Extreme"'],
                               fill=PatternFill("solid", start_color="F47E20")))
    ws["P20"] = "Hand-typed note"
    ws["G13"] = "High"            # someone typed a value for 2 Jul
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace("</worksheet>",
                          X14_EXT % (2 + 1 + len(ratings)) + "</worksheet>")
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    return path


@pytest.fixture
def wb_path(tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "BASE_DIR", str(tmp_path))
    return build_workbook(str(tmp_path / "FDR Daily Statistics Spreadsheet - 2026-27.xlsx"))


def cfg_with(**fdr):
    from app.config import load_config
    c = load_config()
    c["fdr"].update(fdr)
    return c


def test_inspect_finds_the_layout(wb_path):
    lay = workbook.inspect(wb_path)
    # Row 2's summary table names every district too; it has no Date column,
    # so the data header (row 11) is the one used.
    assert lay.sheet == "Sheet1" and lay.header_row == 11
    assert (lay.tfb_col, lay.multi_col, lay.date_col, lay.comment_col) == ("C", "D", "E", "P")
    assert lay.district_cols == COL
    assert lay.rows[date(2026, 7, 1)] == 12
    assert lay.rows[date(2027, 6, 30)] == 376          # followed =E(n-1)+1
    assert lay.rows[D15] == 12 + (D15 - date(2026, 7, 1)).days
    assert lay.allowed["G"][:2] == [PICK, "No Rating"]  # read from the x14 extension
    assert lay.allowed["C"] == ["YES", "NO"] and lay.allowed["D"] == ["YES", "NO"]
    assert lay.placeholders == {PICK.lower()}
    assert lay.tfb_mode == "any"
    assert "dropdown lists stored as Excel extensions" in lay.features


def test_openpyxl_save_would_lose_the_extension_dropdown(wb_path, tmp_path):
    """Documents WHY the writer patches XML instead of saving via openpyxl."""
    out = str(tmp_path / "resaved.xlsx")
    openpyxl.load_workbook(wb_path).save(out)
    with zipfile.ZipFile(out) as z:
        assert b"x14:dataValidation" not in z.read("xl/worksheets/sheet1.xml")


def _row(lay, d):
    return lay.rows[d]


def test_update_writes_confirmed_values_and_preserves_everything(db, wb_path):
    run(when=at(D15, 5, 30))
    run(when=at(D15, 20))
    with zipfile.ZipFile(wb_path) as z:
        before = {n: z.read(n) for n in z.namelist()}
    res = workbook.update_file(wb_path, now=at(D17, 10))
    assert res.ok, res.error
    lay = res.layout
    r15, r16, r17 = _row(lay, D15), _row(lay, D16), _row(lay, D17)
    wb = openpyxl.load_workbook(wb_path)
    ws = wb["Sheet1"]
    assert ws[f"G{r15}"].value == "Extreme"               # replaced the prompt
    assert ws[f"{COL['West and South Gippsland']}{r15}"].value == "Moderate"
    assert ws[f"I{r15}"].value == "High"                   # South West, real column
    assert ws[f"C{r15}"].value == "YES" and ws[f"D{r15}"].value == "NO"
    assert ws[f"C{r16}"].value == "NO" and ws[f"D{r16}"].value == "YES"
    # Central on the 16th had two ratings; the dropdown has no combined value,
    # so the cell is left for a person and the comment says why.
    assert ws[f"{COL['Central']}{r16}"].value == PICK
    assert "Central: High and Extreme" in ws[f"P{r16}"].value
    assert ws[f"P{r16}"].value.startswith("[FDR]")
    # The 17th is "today" but only seen two days ahead: preliminary, unwritten.
    assert ws[f"G{r17}"].value == PICK and ws[f"D{r17}"].value is None
    # Formulas, note, conditional formats, extension untouched.
    assert ws["G3"].value == "=COUNTIF(G12:INDEX(G:G,ROWS(G:G)), C3)"
    assert ws[f"F{r15}"].value == f'=TEXT(E{r15},"ddd")'
    assert ws["D11"].comment.text.startswith("YES/NO")
    assert ws.conditional_formatting
    with zipfile.ZipFile(wb_path) as z:
        after = {n: z.read(n) for n in z.namelist()}
        sheet = z.read("xl/worksheets/sheet1.xml")
        wbxml = z.read("xl/workbook.xml")
    assert b"x14:dataValidation" in sheet
    assert b'fullCalcOnLoad="1"' in wbxml
    changed = {n for n in before if before[n] != after[n]}
    assert changed == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    # Backup of the original exists and is byte-identical to it.
    with zipfile.ZipFile(res.backup) as z:
        assert z.read("xl/worksheets/sheet1.xml") == before["xl/worksheets/sheet1.xml"]
    # Every written cell recorded.
    log = database.read_df("SELECT * FROM fdr_workbook_updates")
    assert len(log) == len(res.plan.changes) and f"G{r15}" in set(log["cell_reference"])


def test_rerun_is_a_no_op(db, wb_path):
    run(when=at(D15, 5, 30))
    assert workbook.update_file(wb_path, now=at(D15, 10)).ok
    with open(wb_path, "rb") as fh:
        first = fh.read()
    res = workbook.update_file(wb_path, now=at(D15, 11))
    assert res.ok and res.plan.changes == []
    with open(wb_path, "rb") as fh:
        assert fh.read() == first


def test_manual_values_are_never_overwritten(db, wb_path):
    run(when=at(D15, 5, 30))
    lay = workbook.inspect(wb_path)
    wb = openpyxl.load_workbook(wb_path)
    wb["Sheet1"][f"G{lay.rows[D15]}"] = "High"   # typed by a person
    wb.save(wb_path)        # (openpyxl save is fine for SETTING UP the test)
    res = workbook.update_file(wb_path, now=at(D15, 10))
    assert any(c == f"G{lay.rows[D15]}" for c, *_ in res.plan.blocked)
    assert openpyxl.load_workbook(wb_path)["Sheet1"][f"G{lay.rows[D15]}"].value == "High"
    assert (store.open_flags()["kind"] == "workbook-blocked").any()
    # --force overrides.
    res = workbook.update_file(wb_path, now=at(D15, 10), force=True)
    assert openpyxl.load_workbook(wb_path)["Sheet1"][f"G{lay.rows[D15]}"].value == "Extreme"


def test_amendment_updates_a_value_this_tool_wrote(db, wb_path):
    run(when=at(D15, 5, 30))
    workbook.update_file(wb_path, now=at(D15, 10))
    amended = FEED.replace(b"Mallee: EXTREME&lt;br&gt;", b"Mallee: CATASTROPHIC&lt;br&gt;", 1)
    run(amended, when=at(D15, 5, 45))     # issued 05:30, before the 06:00 cut: not multi
    res = workbook.update_file(wb_path, now=at(D15, 11))
    r = res.layout.rows[D15]
    assert [c.cell for c in res.plan.changes] == [f"G{r}"]
    assert "amended" in res.plan.changes[0].reason
    assert openpyxl.load_workbook(wb_path)["Sheet1"][f"G{r}"].value == "Catastrophic"


def test_user_comment_is_kept_and_tool_segment_replaced():
    assert workbook.merge_comment("Hand note", "Central: x") == "Hand note | [FDR] Central: x"
    assert workbook.merge_comment("Hand note | [FDR] old", "new") == "Hand note | [FDR] new"
    assert workbook.merge_comment("Hand note | [FDR] old", "") == "Hand note"
    assert workbook.merge_comment(None, "") == ""


def test_combined_value_used_when_the_dropdown_has_one(db, tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "BASE_DIR", str(tmp_path))
    path = build_workbook(str(tmp_path / "wb.xlsx"), ratings=(
        "No Rating", "Moderate", "High", "Extreme", "Catastrophic", "High/Extreme"))
    run(when=at(D15, 20))
    res = workbook.update_file(path, now=at(D16, 10))
    r = res.layout.rows[D16]
    assert openpyxl.load_workbook(path)["Sheet1"][f"L{r}"].value == "High/Extreme"


def test_highest_policy(db, wb_path):
    run(when=at(D15, 20))
    res = workbook.update_file(wb_path, now=at(D16, 10),
                               cfg=cfg_with(multi_rating_display="highest"))
    r = res.layout.rows[D16]
    assert openpyxl.load_workbook(wb_path)["Sheet1"][f"L{r}"].value == "Extreme"


def test_formula_cells_are_never_written(db, tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "BASE_DIR", str(tmp_path))
    path = build_workbook(str(tmp_path / "wb.xlsx"))
    lay = workbook.inspect(path)
    wb = openpyxl.load_workbook(path)
    wb["Sheet1"][f"D{lay.rows[D15]}"] = '=IF(COUNTA(G{0}:O{0})>9,"Yes","No")'.format(
        lay.rows[D15])
    wb.save(path)
    run(when=at(D15, 5, 30))
    res = workbook.update_file(path, now=at(D16, 10))
    assert any(c == f"D{lay.rows[D15]}" and "formula" in why
               for c, _, _, _, why in res.plan.blocked)


def test_dry_run_writes_nothing(db, wb_path):
    run(when=at(D15, 5, 30))
    with open(wb_path, "rb") as fh:
        before = fh.read()
    res = workbook.update_file(wb_path, dry_run=True, now=at(D15, 10))
    assert res.ok and res.plan.changes
    with open(wb_path, "rb") as fh:
        assert fh.read() == before
    assert database.read_df("SELECT COUNT(*) n FROM fdr_workbook_updates").iloc[0]["n"] == 0


def test_open_workbook_is_refused(db, wb_path):
    run(when=at(D15, 5, 30))
    owner = os.path.join(os.path.dirname(wb_path), "~$" + os.path.basename(wb_path))
    open(owner, "w").close()
    res = workbook.update_file(wb_path, now=at(D15, 10))
    assert not res.ok and "open in Excel" in res.error


def test_failed_write_leaves_the_original_untouched(db, wb_path, monkeypatch):
    run(when=at(D15, 5, 30))
    with open(wb_path, "rb") as fh:
        before = fh.read()

    def broken(*a, **k):
        raise workbook.WorkbookError("validation failed (simulated)")
    monkeypatch.setattr(workbook, "validate_written", broken)
    res = workbook.update_file(wb_path, now=at(D15, 10))
    assert not res.ok and "simulated" in res.error
    with open(wb_path, "rb") as fh:
        assert fh.read() == before
    leftovers = [f for f in os.listdir(os.path.dirname(wb_path)) if f.startswith(".fdr-")]
    assert leftovers == []


def test_not_a_workbook(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(workbook.WorkbookError):
        workbook.inspect(str(p))


def test_workbook_without_district_headers(tmp_path):
    p = str(tmp_path / "x.xlsx")
    openpyxl.Workbook().save(p)
    with pytest.raises(workbook.WorkbookError, match="district"):
        workbook.inspect(p)


def test_duplicate_dates_are_noted(tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "BASE_DIR", str(tmp_path))
    path = build_workbook(str(tmp_path / "wb.xlsx"))
    wb = openpyxl.load_workbook(path)
    wb["Sheet1"]["E14"] = datetime(2026, 7, 1)   # a mistyped date
    wb.save(path)
    lay = workbook.inspect(path)
    assert any("appears on rows" in n for n in lay.notes)


def test_upload_fill_bytes(db, wb_path):
    run(when=at(D15, 5, 30))
    with open(wb_path, "rb") as fh:
        data = fh.read()
    res, out = workbook.fill_bytes(data, "FDR.xlsx", now=at(D15, 10))
    assert res.ok and out and out != data
    res, out = workbook.fill_bytes(b"x", "notes.txt")
    assert not res.ok and out is None


# --------------------------------------------------------------------------- #
# CLI + page access
# --------------------------------------------------------------------------- #
def test_cli_fetch_from_file_and_dry_run(db, wb_path, tmp_path, monkeypatch, capsys):
    from app.modules.fdr import cli
    monkeypatch.setattr(cli, "LOCK_FILE", str(tmp_path / "fdr.lock"))
    monkeypatch.setattr(cli, "LOG_FILE", str(tmp_path / "fdr.log"))
    monkeypatch.setattr(model, "now", lambda: at(D15, 5, 30))
    feed = tmp_path / "feed.xml"
    feed.write_bytes(FEED)
    assert cli.main(["fetch", "--from-file", str(feed)]) == 0
    code = cli.main(["update-excel", "--workbook", wb_path, "--dry-run"])
    out = capsys.readouterr().out
    assert "Would write" in out and code in (0, 2)


def test_cli_refuses_a_second_concurrent_run(db, tmp_path, monkeypatch):
    from app.modules.fdr import cli
    lock = tmp_path / "fdr.lock"
    lock.write_text("1 0")
    monkeypatch.setattr(cli, "LOCK_FILE", str(lock))
    monkeypatch.setattr(cli, "LOG_FILE", str(tmp_path / "fdr.log"))
    with pytest.raises(SystemExit) as e:
        cli.main(["fetch", "--from-file", str(lock)])
    assert e.value.code == 3


def test_fdr_callbacks_need_the_intel_password(db):
    from app.factory import create_app
    client = create_app(autostart=False).server.test_client()
    r = client.post("/_dash-update-component", json={
        "output": "..fdr-status.children...fdr-forecast.children...fdr-fetch-msg.children..",
        "outputs": [{"id": "fdr-status", "property": "children"},
                    {"id": "fdr-forecast", "property": "children"},
                    {"id": "fdr-fetch-msg", "property": "children"}],
        "inputs": [{"id": "fdr-tick", "property": "n_intervals", "value": 1},
                   {"id": "fdr-fetch", "property": "n_clicks", "value": 1}],
        "changedPropIds": ["fdr-fetch.n_clicks"], "state": []})
    assert r.status_code == 204           # PreventUpdate: nothing ran
