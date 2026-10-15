# Passive Monitor — Copyright (c) 2026 SirTophamMatt. All rights reserved.
"""The FDR Daily Statistics workbook: discover its layout, plan, write safely.

Why the write does NOT go through openpyxl
------------------------------------------
openpyxl rebuilds the whole file on save and silently drops what it does not
model: list validations stored as Excel extensions (``x14:dataValidation`` —
what Excel writes when a dropdown's list lives on another sheet), extension
conditional formatting (data bars/icon sets), images and charts it cannot
round-trip, slicers, sparklines, threaded comments, and custom XML. On a
statistics workbook that is exactly the formatting, dropdowns and summaries the
user needs kept. So:

* openpyxl is used READ-ONLY, to find the sheet, the header row, the date of
  every row, the existing values, the dropdown lists and every formula;
* the write opens the .xlsx (a zip), edits ONLY the ``<c>`` elements of the
  target cells in that one worksheet's XML, sets ``fullCalcOnLoad`` so Excel
  recalculates the summary formulas on open, and copies every other part
  byte-for-byte.

Safety sequence (``update_file``): refuse if Excel has the file open ->
inspect -> plan -> write to a temp file beside the original -> validate the
temp file (re-open, read back every changed cell, same sheets, same formula
count, untouched parts byte-identical) -> back up the original -> atomic
replace. Any failure leaves the original untouched.

Write rules (``plan``)
----------------------
* rows are matched by their DATE value, never by row number;
* a formula cell is never written;
* an empty cell is filled; a cell holding the value this tool last wrote there
  is updated (an amendment); a cell holding anything else was typed by a
  person and is left alone and flagged (``force`` overrides);
* days that have started get CONFIRMED values only (see ``resolve``);
  unknown stays blank. A rating that changed during the day records the
  LATEST rating, with what it was and when it changed in Comments;
* days still to come get CFA's current forecast marked "(forecast)" plus a
  dated Comments note — provisional, replaced when the day starts, not logged.
"""
import io
import logging
import os
import re
import shutil
import tempfile
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter

from app.config import BASE_DIR, load_config
from app.modules.fdr import model, resolve, store

log = logging.getLogger(__name__)

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_XM = "http://schemas.microsoft.com/office/excel/2006/main"
NS_X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
TAG = "[FDR]"
_CELL_RE = re.compile(r"^([A-Z]+)(\d+)$")

# Zip parts whose presence is reported (all preserved untouched by the writer).
_FEATURES = [
    (r"xl/vbaProject\.bin", "VBA macros"),
    (r"xl/charts/", "charts"),
    (r"xl/drawings/", "drawings / images / shapes"),
    (r"xl/pivotTables/", "pivot tables"),
    (r"xl/externalLinks/", "links to other workbooks"),
    (r"xl/tables/", "Excel tables"),
    (r"xl/slicers/", "slicers"),
    (r"xl/threadedComments/", "threaded comments"),
    (r"xl/comments\d*\.xml", "cell notes"),
    (r"customXml/", "custom XML (e.g. SharePoint/sensitivity labels)"),
]


class WorkbookError(Exception):
    pass


class WorkbookLocked(WorkbookError):
    pass


@dataclass
class Layout:
    path: str
    sheet: str
    sheet_part: str
    header_row: int
    date_col: str
    district_cols: dict          # district -> column letter
    tfb_col: str = None
    multi_col: str = None
    comment_col: str = None
    rows: dict = field(default_factory=dict)        # date -> row number
    values: dict = field(default_factory=dict)      # "G12" -> current value
    formulas: set = field(default_factory=set)      # cells holding a formula
    allowed: dict = field(default_factory=dict)     # column -> [allowed values]
    # Dropdown entries that are prompts, not values ("Pick Rating…."): a cell
    # holding one is treated as empty.
    placeholders: set = field(default_factory=set)
    existing: dict = field(default_factory=dict)    # column -> {values seen}
    formula_count: int = 0
    sheet_names: list = field(default_factory=list)
    features: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    tfb_mode: str = "any"

    @property
    def first_date(self):
        return min(self.rows) if self.rows else None

    @property
    def last_date(self):
        return max(self.rows) if self.rows else None

    def describe(self):
        lines = [f"Sheet '{self.sheet}', header row {self.header_row}, "
                 f"{len(self.rows)} dated rows "
                 f"({self.first_date} .. {self.last_date})",
                 f"Date column {self.date_col}; TFB {self.tfb_col or '—'} "
                 f"(convention: {self.tfb_mode}); Multi Day {self.multi_col or '—'}; "
                 f"Comments {self.comment_col or '—'}",
                 "Districts: " + ", ".join(f"{d} {c}" for d, c in self.district_cols.items())]
        if self.features:
            lines.append("Present and preserved untouched: " + ", ".join(self.features))
        lines += self.notes
        return lines


@dataclass
class Change:
    cell: str
    valid_date: date
    old: object
    new: object
    reason: str
    forecast: bool = False     # a future day's provisional value: not logged


@dataclass
class Plan:
    changes: list = field(default_factory=list)
    blocked: list = field(default_factory=list)     # (cell, date, old, new, why)
    review: list = field(default_factory=list)      # (date, note)
    no_data: list = field(default_factory=list)     # started days with nothing stored

    def no_data_summary(self):
        if not self.no_data:
            return None
        return (f"{len(self.no_data)} started day(s) have no stored data "
                f"({self.no_data[0]:%d/%m/%Y} .. {self.no_data[-1]:%d/%m/%Y}); their "
                "cells are left exactly as they are — CFA publishes no archive, so "
                "only days collected by this tool (or corrected here) can be filled")


@dataclass
class Result:
    ok: bool
    dry_run: bool
    path: str
    layout: Layout = None
    plan: Plan = None
    backup: str = None
    error: str = None
    run_id: int = None

    def summary(self):
        p = self.plan or Plan()
        return {"ok": self.ok, "dry_run": self.dry_run, "path": self.path,
                "changes": len(p.changes), "blocked": len(p.blocked),
                "review": len(p.review), "backup": self.backup, "error": self.error}


# --------------------------------------------------------------------------- #
# Discovery
# --------------------------------------------------------------------------- #
def _load(path, data_only):
    """openpyxl, read-only use. Its "Data Validation extension is not supported"
    warning is expected: extension dropdowns are read from the XML here and are
    never lost, because nothing is ever saved through openpyxl."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*extension is not supported.*")
        return openpyxl.load_workbook(path, data_only=data_only)


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d %B %Y", "%d %b %Y",
                    "%A, %d %B %Y", "%a %d/%m/%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _sheet_parts(zf):
    """{sheet name: zip part path}."""
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for s in wb.iter(f"{{{NS}}}sheet"):
        t = targets.get(s.get(f"{{{NS_R}}}id"), "")
        t = t.lstrip("/")
        out[s.get("name")] = t if t.startswith("xl/") else "xl/" + t
    return out


def _cells_in(sqref):
    """Column letters covered by an sqref ("G12:O376 C12:C376")."""
    cols = set()
    for part in (sqref or "").split():
        a, _, b = part.partition(":")
        ma, mb = re.match(r"\$?([A-Z]+)", a), re.match(r"\$?([A-Z]+)", b or a)
        if ma and mb:
            for i in range(column_index_from_string(ma.group(1)),
                           column_index_from_string(mb.group(1)) + 1):
                cols.add(get_column_letter(i))
    return cols


def _list_values(formula, wb_values):
    """Allowed values of a list validation formula (inline or a range)."""
    f = (formula or "").strip()
    if f.startswith('"') and f.endswith('"'):
        return [x.strip() for x in f[1:-1].split(",") if x.strip()]
    f = f.lstrip("=")
    defined = wb_values.defined_names.get(f) if hasattr(
        wb_values.defined_names, "get") else None
    if defined is not None:
        f = defined.attr_text
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", f)
    if not m or m.group(1) not in wb_values.sheetnames:
        return None
    ws = wb_values[m.group(1)]
    c1, r1 = m.group(2), int(m.group(3))
    c2, r2 = m.group(4) or c1, int(m.group(5) or r1)
    out = []
    for row in ws.iter_rows(min_row=r1, max_row=r2,
                            min_col=column_index_from_string(c1),
                            max_col=column_index_from_string(c2), values_only=True):
        out += [_text(v) for v in row if _text(v)]
    return out


def _validations(sheet_xml, wb_values):
    """{column: [allowed values]} from standard AND x14 list validations."""
    root = etree.fromstring(sheet_xml)
    out = {}
    for dv in root.iter(f"{{{NS}}}dataValidation"):
        if dv.get("type") != "list":
            continue
        vals = _list_values(dv.findtext(f"{{{NS}}}formula1"), wb_values)
        for col in _cells_in(dv.get("sqref")):
            if vals:
                out[col] = vals
    for dv in root.iter(f"{{{NS_X14}}}dataValidation"):
        if dv.get("type") != "list":
            continue
        f = dv.find(f"{{{NS_X14}}}formula1")
        vals = _list_values(f.findtext(f"{{{NS_XM}}}f") if f is not None else None,
                            wb_values)
        for col in _cells_in(dv.findtext(f"{{{NS_XM}}}sqref")):
            if vals:
                out[col] = vals
    return out


def _find_header(ws, max_rows=40):
    """(header row, {district: col}) — the first row naming >= 7 districts
    that also has a Date column. A summary table above the data (the FDR
    workbook's rows 2–8: district totals per rating) names the districts too,
    but has no Date, so it is skipped."""
    for row in ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row)):
        found = {}
        for c in row:
            d = model.normalise_district(_text(c.value))
            if d and d not in found:
                found[d] = c.column_letter
        if len(found) >= 7 and "date" in _label_cols(ws, row[0].row):
            return row[0].row, found
    return None, {}


def _label_cols(ws, header_row):
    """Date / TFB / Multi Day / Comments columns from the header row and the
    two rows above it (headers are often merged over two rows)."""
    out = {}
    for r in range(max(1, header_row - 2), header_row + 1):
        for c in ws[r]:
            t = _text(c.value).lower()
            if not t:
                continue
            key = None
            if re.search(r"\bmulti", t):
                key = "multi"
            elif re.search(r"\btfb\b|total\s*fire\s*ban", t):
                key = "tfb"
            elif re.search(r"\bcomment|\bnotes?\b|\bremarks?\b", t):
                key = "comment"
            elif re.fullmatch(r"date|day|date of", t) or t.startswith("date"):
                key = "date"
            if key and key not in out:
                out[key] = c.column_letter
    return out


def inspect(path, cfg=None):
    """Read the workbook and return its Layout. Raises WorkbookError."""
    cfg = (cfg or load_config())["fdr"]
    if not os.path.exists(path):
        raise WorkbookError(f"Workbook not found: {path}")
    try:
        wb_v = _load(path, True)
        wb_f = _load(path, False)
    except Exception as e:
        raise WorkbookError(f"Not a readable .xlsx/.xlsm workbook: {e}") from e
    sheet = cfg.get("workbook_sheet") or None
    candidates = [sheet] if sheet else wb_v.sheetnames
    header_row = districts = None
    for name in candidates:
        if name not in wb_v.sheetnames:
            raise WorkbookError(f"Sheet {name!r} not in workbook")
        header_row, districts = _find_header(wb_v[name])
        if header_row:
            sheet = name
            break
    if not header_row:
        raise WorkbookError("No header row naming the fire districts was found "
                            "(looked for at least 7 of the 9 district names in "
                            "the first 40 rows of each sheet).")
    ws_v, ws_f = wb_v[sheet], wb_f[sheet]
    labels = _label_cols(ws_v, header_row)
    if "date" not in labels:
        raise WorkbookError(f"No 'Date' column found near header row {header_row}.")
    with zipfile.ZipFile(path) as zf:
        part = _sheet_parts(zf)[sheet]
        names = zf.namelist()
        allowed = _validations(zf.read(part), wb_v)
        has_x14 = b"x14:dataValidation" in zf.read(part)
    lay = Layout(path=path, sheet=sheet, sheet_part=part, header_row=header_row,
                 date_col=labels["date"], district_cols=districts,
                 tfb_col=labels.get("tfb"), multi_col=labels.get("multi"),
                 comment_col=labels.get("comment"), allowed=allowed,
                 sheet_names=list(wb_v.sheetnames))
    missing = [d for d in model.DISTRICTS if d not in districts]
    if missing:
        lay.notes.append("Districts with no column: " + ", ".join(missing))
    for pat, label in _FEATURES:
        if any(re.match(pat, n) for n in names):
            lay.features.append(label)
    if has_x14:
        lay.features.append("dropdown lists stored as Excel extensions")
    # Every formula in the workbook (validation compares this count after write).
    for ws in wb_f.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    lay.formula_count += 1
    targets = [lay.date_col, *districts.values()] + [
        c for c in (lay.tfb_col, lay.multi_col, lay.comment_col) if c]
    prev = None
    for r in range(header_row + 1, ws_v.max_row + 1):
        for col in targets:
            ref = f"{col}{r}"
            fv = ws_f[ref].value
            if isinstance(fv, str) and fv.startswith("="):
                lay.formulas.add(ref)
            v = ws_v[ref].value
            if v is not None and _text(v) != "":
                lay.values[ref] = v
                if col != lay.date_col:
                    lay.existing.setdefault(col, set()).add(_text(v))
        d = _as_date(ws_v[f"{lay.date_col}{r}"].value)
        fv = ws_f[f"{lay.date_col}{r}"].value
        if d is None and prev and isinstance(fv, str):
            m = re.match(r"^=\$?[A-Z]+\$?(\d+)\s*\+\s*1$", fv.replace(" ", ""))
            if m and int(m.group(1)) == r - 1:
                d = prev + timedelta(days=1)   # =E12+1 with no cached value
        if d is None:
            continue
        if d in lay.rows:
            lay.notes.append(f"Date {d} appears on rows {lay.rows[d]} and {r}; "
                             "only the first is used")
            continue
        lay.rows[d] = r
        prev = d
    if not lay.rows:
        raise WorkbookError(f"No dates found in column {lay.date_col} below row "
                            f"{header_row}.")
    for vals in allowed.values():
        for v in vals:
            if not model.parse_ratings(v)[1] and v.lower() not in (
                    "yes", "no", "y", "n", "true", "false"):
                lay.placeholders.add(v.lower())
    # Prompts and this tool's own forecast values say nothing about how the
    # workbook spells a real value.
    ftag = (cfg.get("forecast_suffix") or "").strip().lower()
    for col in list(lay.existing):
        lay.existing[col] = {v for v in lay.existing[col]
                             if v.lower() not in lay.placeholders
                             and not (ftag and v.lower().endswith(ftag))}
    lay.tfb_mode = _tfb_mode(lay, cfg.get("tfb_mode", "auto"))
    return lay


def _tfb_mode(lay, configured):
    if configured and configured != "auto":
        return configured
    if not lay.tfb_col:
        return "none"
    vals = list(lay.allowed.get(lay.tfb_col, [])) + list(lay.existing.get(lay.tfb_col, []))
    if any(model.districts_in(v) for v in vals):
        return "districts"
    if vals and all(re.fullmatch(r"\d+", v) for v in vals):
        return "count"
    return "any"


# --------------------------------------------------------------------------- #
# Conventions — turn a canonical value into the workbook's own spelling
# --------------------------------------------------------------------------- #
def _vocab(lay, col):
    return list(lay.allowed.get(col, [])) + sorted(lay.existing.get(col, set()))


def rating_value(lay, col, ratings, policy="auto", sep="/"):
    """(value | None, note). None = do not write."""
    vocab = _vocab(lay, col)
    parsed = {}
    for v in vocab:
        found, ok = model.parse_ratings(v)
        if ok and found:
            parsed.setdefault(frozenset(found), v)
    key = frozenset(ratings)
    if len(ratings) == 1:
        if key in parsed:
            return parsed[key], None
        if lay.allowed.get(col):
            return None, (f"{ratings[0]} is not in the column's dropdown list "
                          f"({', '.join(lay.allowed[col])})")
        style = next(iter(vocab), "")
        r = ratings[0]
        return (r.upper() if style.isupper() else r), None
    # More than one rating in the district on the day.
    if policy in ("auto", "combined") and key in parsed:
        return parsed[key], None
    if policy == "highest":
        return rating_value(lay, col, [model.highest(ratings)], "auto", sep)
    if policy == "combined" and not lay.allowed.get(col):
        style = next(iter(vocab), "")
        txt = model.display(ratings, sep)
        return (txt.upper() if style.isupper() else txt), None
    return None, ("more than one rating — left blank for manual entry "
                  "(the workbook has no matching combined value; set "
                  "fdr.multi_rating_display to 'highest' or 'combined' to change)")


def yes_no_value(lay, col, yes):
    vocab = _vocab(lay, col)
    want = ("yes", "y", "true") if yes else ("no", "n", "false")
    for v in vocab:
        if v.lower() in want:
            return v
    style = next(iter(vocab), "")
    word = "Yes" if yes else "No"
    return word.upper() if style.isupper() else word


def merge_comment(existing, ours):
    """Keep a person's comment; replace only this tool's tagged segment."""
    base = re.sub(r"\s*(\|\s*)?" + re.escape(TAG) + r".*$", "", _text(existing),
                  flags=re.S).rstrip(" |")
    if not ours:
        return base
    return f"{base} | {TAG} {ours}" if base else f"{TAG} {ours}"


# --------------------------------------------------------------------------- #
# Plan
# --------------------------------------------------------------------------- #
def _same(a, b):
    return _text(a).lower() == _text(b).lower()


def plan(lay, records, written, *, policy="latest", sep="/", force=False,
         forecasts=True, suffix=" (forecast)"):
    """Plan the cell changes for ``records`` ({date: DayRecord}).

    Days that have started get their confirmed values (logged in
    fdr_workbook_updates). Days still to come get CFA's current forecast with
    ``suffix`` on every value ("Moderate (forecast)") and a dated note in
    Comments: visibly provisional, coloured by the sheet's own SEARCH-based
    formatting, NOT counted by the summary's exact-match COUNTIFs, and not
    logged. A forecast cell is the tool's to replace — until the day starts
    and the confirmed value takes its place."""
    p = Plan()
    tag = suffix.strip().lower()

    def is_forecast(value):
        return bool(tag) and _text(value).lower().endswith(tag)

    def want(cell, d, new, reason, forecast=False):
        if new is None:
            return
        old = lay.values.get(cell)
        if cell in lay.formulas:
            if not forecast:
                p.blocked.append((cell, d, old, new, "cell holds a formula"))
            return
        if _same(old, new):
            return
        if _text(old) == "" or _text(old).lower() in lay.placeholders \
                or is_forecast(old):
            p.changes.append(Change(cell, d, old, new, reason, forecast))
        elif force or _same(written.get((lay.sheet, cell)), old):
            p.changes.append(Change(cell, d, old, new,
                                    reason + (" (forced)" if force else " (amended)")))
        elif not forecast:
            # A person's entry. For a future day it is left silently — the
            # forecast is only provisional anyway.
            p.blocked.append((cell, d, old, new,
                              "holds a value this tool did not write — left as is"))

    def comment(d, r, ours):
        if not lay.comment_col:
            return
        cell = f"{lay.comment_col}{r}"
        old = lay.values.get(cell)
        new = merge_comment(old, ours)
        if not _same(old, new) and cell not in lay.formulas:
            # Comments are merged, never overwritten: a person's text is kept
            # and only this tool's [FDR] segment is replaced.
            p.changes.append(Change(cell, d, old, new or None, "Comments",
                                    is_forecast_day))

    for d, rec in sorted(records.items()):
        r = lay.rows.get(d)
        if r is None:
            continue
        is_forecast_day = rec.stage == "future"
        if is_forecast_day:
            if forecasts and not rec.no_data:
                _plan_forecast(lay, rec, d, r, want, comment, sep, suffix)
            continue
        if rec.no_data:
            p.no_data.append(d)
            continue
        notes = []
        for dist, col in lay.district_cols.items():
            dd = rec.districts[dist]
            if dd.state != "confirmed" or not dd.ratings:
                continue
            ratings, pol = dd.ratings, policy
            if policy == "latest":
                # The rating changed during the day: record the one in force
                # at the end (Comments says what it was before, and when).
                # Two ratings published at once have no "latest" — they fall
                # back to the workbook's own combined value or review.
                if dd.multi and len(dd.latest) == 1:
                    ratings = dd.latest
                pol = "auto"
            val, note = rating_value(lay, col, ratings, pol, sep)
            if note:
                notes.append(f"{dist}: {note}")
                p.review.append((d, f"{dist} ({col}{r}): {note}"))
            want(f"{col}{r}", d, val, f"{dist} rating ({dd.source})")
        # Multi Day: "did the FDR change during the day". YES as soon as a
        # change is confirmed (it cannot un-happen); NO only once the day is
        # over and nothing changed.
        if lay.multi_col and (rec.multi == "YES" or
                              (rec.multi == "NO" and rec.stage == "past")):
            want(f"{lay.multi_col}{r}", d, yes_no_value(lay, lay.multi_col,
                                                        rec.multi == "YES"),
                 f"Multi Day ({rec.multi_source})")
        if lay.tfb_col and rec.tfb_any:
            want(f"{lay.tfb_col}{r}", d, _tfb_value(lay, rec, rec.tfb_any),
                 "Total Fire Ban")
        ours = rec.comment()
        blank = [n.split(":")[0] for n in notes if "left blank" in n]
        if blank:
            ours = (ours + "; " if ours else "") + \
                "rating left blank for: " + ", ".join(blank)
        comment(d, r, ours)
        for n in rec.review:
            p.review.append((d, n))
    return p


def _tfb_value(lay, rec, status, districts=None):
    """The TFB cell value. The FDR workbook's column asks "is a Total Fire Ban
    declared?" — YES if any district has one (mode ``any``)."""
    districts = rec.tfb_districts if districts is None else districts
    if lay.tfb_mode == "any":
        return yes_no_value(lay, lay.tfb_col, status == "YES")
    if lay.tfb_mode == "districts":
        return ", ".join(districts) if status == "YES" else None
    if lay.tfb_mode == "count":
        return len(districts) if status == "YES" else 0
    return None


def _plan_forecast(lay, rec, d, r, want, comment, sep, suffix):
    """A day not yet started: CFA's current forecast, marked as such."""
    wrote = False
    for dist, col in lay.district_cols.items():
        dd = rec.districts[dist]
        if dd.state in ("conflict", "unrecognised") or len(dd.latest) != 1:
            continue
        val, _ = rating_value(lay, col, dd.latest, "auto", sep)
        if val:
            want(f"{col}{r}", d, val + suffix, f"{dist} forecast", forecast=True)
            wrote = True
    # A day CFA has no issued forecast for yet (its placeholder fifth day)
    # gets nothing — a TFB "forecast" alone would look like more than it is.
    if wrote and lay.tfb_col and rec.tfb_forecast and lay.tfb_mode in ("any", "count"):
        yes = [x.district for x in rec.districts.values() if x.tfb == model.TFB_YES]
        v = _tfb_value(lay, rec, rec.tfb_forecast, yes)
        want(f"{lay.tfb_col}{r}", d, f"{v}{suffix}", "TFB forecast", forecast=True)
        wrote = True
    if wrote:
        issued = rec.forecast_issued
        when = (datetime.fromisoformat(issued).strftime("%d/%m %H:%M")
                if issued else None)
        comment(d, r, "Forecast" + (f" (BoM issued {when})" if when else "")
                + " — future day, may change")


# --------------------------------------------------------------------------- #
# Writer — direct XML patch of one worksheet part
# --------------------------------------------------------------------------- #
def _col_index(ref):
    return column_index_from_string(_CELL_RE.match(ref).group(1))


def _set_cell(sheet_data, ref, value, col_styles):
    col, row = _CELL_RE.match(ref).groups()
    row = int(row)
    rows = sheet_data.findall(f"{{{NS}}}row")
    row_el = next((x for x in rows if int(x.get("r")) == row), None)
    if row_el is None:
        row_el = etree.Element(f"{{{NS}}}row", r=str(row))
        after = [x for x in rows if int(x.get("r")) < row]
        if after:
            after[-1].addnext(row_el)
        else:
            sheet_data.insert(0, row_el)
    cells = row_el.findall(f"{{{NS}}}c")
    c = next((x for x in cells if x.get("r") == ref), None)
    if c is None:
        c = etree.Element(f"{{{NS}}}c", r=ref)
        style = row_el.get("s") if row_el.get("customFormat") == "1" else \
            col_styles.get(column_index_from_string(col))
        if style:
            c.set("s", style)
        before = [x for x in cells if _col_index(x.get("r")) < column_index_from_string(col)]
        if before:
            before[-1].addnext(c)
        else:
            row_el.insert(0, c)
    if c.find(f"{{{NS}}}f") is not None:
        raise WorkbookError(f"{ref} holds a formula; refusing to overwrite it")
    for child in list(c):
        c.remove(child)
    if value is None or value == "":
        c.attrib.pop("t", None)
        return
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        c.attrib.pop("t", None)
        etree.SubElement(c, f"{{{NS}}}v").text = repr(value) if isinstance(
            value, float) else str(value)
        return
    c.set("t", "inlineStr")
    is_ = etree.SubElement(c, f"{{{NS}}}is")
    t = etree.SubElement(is_, f"{{{NS}}}t")
    t.text = str(value)
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _col_styles(root):
    out = {}
    for col in root.iter(f"{{{NS}}}col"):
        if col.get("style"):
            for i in range(int(col.get("min")), int(col.get("max")) + 1):
                out[i] = col.get("style")
    return out


def _serialise(root):
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                          standalone=True)


def write_changes(src, dst, sheet_part, changes):
    """Copy ``src`` to ``dst`` with ``changes`` applied to ``sheet_part``."""
    with zipfile.ZipFile(src) as zin:
        root = etree.fromstring(zin.read(sheet_part),
                                etree.XMLParser(huge_tree=True, remove_blank_text=False))
        data = root.find(f"{{{NS}}}sheetData")
        styles = _col_styles(root)
        for ch in changes:
            _set_cell(data, ch.cell, ch.new, styles)
        new_sheet = _serialise(root)
        wb = etree.fromstring(zin.read("xl/workbook.xml"))
        calc = wb.find(f"{{{NS}}}calcPr")
        if calc is None:
            calc = etree.Element(f"{{{NS}}}calcPr")
            anchor = wb.find(f"{{{NS}}}definedNames")
            if anchor is None:
                anchor = wb.find(f"{{{NS}}}sheets")
            anchor.addnext(calc)
        calc.set("fullCalcOnLoad", "1")
        new_wb = _serialise(wb)
        with zipfile.ZipFile(dst, "w") as zout:
            for info in zin.infolist():
                if info.filename == sheet_part:
                    zout.writestr(info, new_sheet, compress_type=zipfile.ZIP_DEFLATED)
                elif info.filename == "xl/workbook.xml":
                    zout.writestr(info, new_wb, compress_type=zipfile.ZIP_DEFLATED)
                else:
                    zout.writestr(info, zin.read(info.filename),
                                  compress_type=info.compress_type)


def validate_written(src, dst, lay, changes):
    """Raise WorkbookError unless ``dst`` is ``src`` plus exactly ``changes``."""
    with zipfile.ZipFile(dst) as z2:
        bad = z2.testzip()
        if bad:
            raise WorkbookError(f"written file is corrupt at {bad}")
        with zipfile.ZipFile(src) as z1:
            if z1.namelist() != z2.namelist():
                raise WorkbookError("written file does not have the same parts")
            for n in z1.namelist():
                if n not in (lay.sheet_part, "xl/workbook.xml") and \
                        z1.read(n) != z2.read(n):
                    raise WorkbookError(f"part {n} changed unexpectedly")
    try:
        wb = _load(dst, False)
    except Exception as e:
        raise WorkbookError(f"written file does not open: {e}") from e
    if wb.sheetnames != lay.sheet_names:
        raise WorkbookError("sheet list changed")
    count = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("="))
    if count != lay.formula_count:
        raise WorkbookError(f"formula count changed ({lay.formula_count} -> {count})")
    ws = wb[lay.sheet]
    for ch in changes:
        got = ws[ch.cell].value
        if _text(got) != _text(ch.new):
            raise WorkbookError(f"{ch.cell} reads back {got!r}, expected {ch.new!r}")


# --------------------------------------------------------------------------- #
# Orchestration
# --------------------------------------------------------------------------- #
def check_not_open(path):
    """Raise WorkbookLocked if Excel (or anything else) has the file open."""
    folder, name = os.path.split(os.path.abspath(path))
    # Excel's owner file is "~$" + the name, with the first two characters of
    # the name dropped when it is long; check both spellings.
    for cand in (os.path.join(folder, "~$" + name),
                 os.path.join(folder, "~$" + name[2:])):
        if os.path.exists(cand):
            raise WorkbookLocked(f"{name} appears to be open in Excel (lock file "
                                 f"{os.path.basename(cand)} exists). Close it and re-run.")
    try:
        with open(path, "r+b"):
            pass
    except PermissionError as e:
        raise WorkbookLocked(f"{name} is locked or read-only: {e}") from e


def backup_dir(cfg):
    d = cfg.get("backup_dir") or os.path.join(BASE_DIR, "fdr_backups")
    os.makedirs(d, exist_ok=True)
    return d


def _backup(path, cfg):
    d = backup_dir(cfg)
    stem, ext = os.path.splitext(os.path.basename(path))
    dest = os.path.join(d, f"{stem}.{model.now().strftime('%Y%m%d_%H%M%S')}{ext}")
    shutil.copy2(path, dest)
    keep = int(cfg.get("backups_keep", 30))
    olds = sorted(f for f in os.listdir(d) if f.startswith(stem + ".") and f.endswith(ext))
    for f in olds[:-keep]:
        try:
            os.remove(os.path.join(d, f))
        except OSError:
            pass
    return dest


def records_for(lay, start=None, end=None, cfg=None, now=None):
    now = now or model.now()
    lo = max(lay.first_date, model.as_date(start)) if start else lay.first_date
    hi = min(lay.last_date, model.as_date(end)) if end else lay.last_date
    if lo > hi:
        return {}
    return resolve.resolve_range(lo, hi, cfg=cfg, now=now)


def update_file(path, *, run_id=None, cfg=None, dry_run=False, force=False,
                start=None, end=None, output=None, workbook_name=None, now=None):
    """Update a workbook in place (or into ``output``). Returns a Result.

    Never raises for an expected failure (locked, unreadable, validation) —
    the Result says what happened and the original is untouched."""
    full = cfg or load_config()
    c = full["fdr"]
    name = workbook_name or os.path.basename(path)
    res = Result(ok=False, dry_run=dry_run, path=output or path)
    try:
        if not dry_run and not output:
            check_not_open(path)
        lay = inspect(path, full)
        res.layout = lay
        recs = records_for(lay, start, end, full, now)
        written = store.written_cells(name)
        res.plan = plan(lay, recs, written,
                        policy=c.get("multi_rating_display", "latest"),
                        sep=c.get("combined_separator", "/"), force=force,
                        forecasts=c.get("forecast_days", True),
                        suffix=c.get("forecast_suffix", " (forecast)"))
        if dry_run or not res.plan.changes:
            res.ok = True
            _flag_blocked(res.plan, name)
            return res
        target = output or path
        fd, tmp = tempfile.mkstemp(prefix=".fdr-", suffix=os.path.splitext(path)[1],
                                   dir=os.path.dirname(os.path.abspath(target)))
        os.close(fd)
        try:
            write_changes(path, tmp, lay.sheet_part, res.plan.changes)
            validate_written(path, tmp, lay, res.plan.changes)
            if not output:
                res.backup = _backup(path, c)
            os.replace(tmp, target)
        except Exception:
            if os.path.exists(tmp):
                os.remove(tmp)
            raise
        if run_id is None:
            run_id = store.start_run("workbook")
            store.finish_run(run_id, "ok", {"updated": len(res.plan.changes)}, [])
        res.run_id = run_id
        stamp = model.iso(model.now())
        # Forecast cells (future days) are provisional and not logged.
        store.record_cell_updates([{
            "run_id": run_id, "workbook_name": name, "worksheet": lay.sheet,
            "cell_reference": ch.cell, "valid_date": ch.valid_date.isoformat(),
            "previous_value": None if ch.old is None else _text(ch.old),
            "new_value": None if ch.new is None else _text(ch.new),
            "reason": ch.reason, "dry_run": 0, "updated_at": stamp}
            for ch in res.plan.changes if not ch.forecast])
        _flag_blocked(res.plan, name)
        res.ok = True
    except WorkbookLocked as e:
        res.error = str(e)
    except WorkbookError as e:
        res.error = str(e)
    except PermissionError as e:
        res.error = f"Could not replace the workbook (is it open?): {e}"
    except Exception as e:
        log.exception("FDR workbook update failed")
        res.error = f"Unexpected error: {e}"
    if res.error:
        log.warning("FDR workbook %s: %s", name, res.error)
    return res


def _flag_blocked(p, name):
    if not p.blocked:
        return
    conn = store.connect()
    try:
        for cell, d, old, new, why in p.blocked:
            store.add_flag(conn, d.isoformat(), None, "workbook-blocked",
                           f"{name} {cell}: would write {new!r} over {old!r} — {why}")
        conn.commit()
    finally:
        conn.close()


def fill_bytes(data, filename, **kw):
    """Uploaded workbook bytes -> (Result, filled bytes | None). The upload
    is the 'original'; nothing on the server is replaced."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in (".xlsx", ".xlsm"):
        r = Result(ok=False, dry_run=kw.get("dry_run", False), path=filename)
        r.error = "Upload an .xlsx or .xlsm workbook."
        return r, None
    with tempfile.TemporaryDirectory(prefix="fdr-up-") as tmp:
        src = os.path.join(tmp, "in" + ext)
        out = os.path.join(tmp, "out" + ext)
        with open(src, "wb") as fh:
            fh.write(data)
        res = update_file(src, output=None if kw.get("dry_run") else out,
                          workbook_name=filename, **kw)
        res.path = filename
        if not res.ok or kw.get("dry_run"):
            return res, None
        if not res.plan.changes:
            return res, data
        with open(out, "rb") as fh:
            return res, fh.read()
