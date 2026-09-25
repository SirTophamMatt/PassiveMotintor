"""Render an Operational Summary draft into the SCC's own PowerPoint template.

The output is the template deck (``seed/opsum_template.pptx``, built from an
issued summary by ``seed/opsum_template_tool.py``) with its tokens filled in,
not a re-drawing of it — fonts, colours, menu links and icons are the SCC's.

What gets filled (see ``app/opsum.py`` for the field map):

* A paragraph that is exactly ``{{key}}`` for a *text* field expands into one
  paragraph per line, each a copy of the token paragraph (same bullet, indent,
  spacing, font). Sub-bullets, headings, spacers, bold and links come from the
  line syntax documented in ``opsum``.
* Any other ``{{token}}`` is replaced inside its run (``**bold**`` splits the
  run so the bold part keeps every other font property).
* Cells of the AV workload grid are recoloured and blanked; the flood Status
  row is coloured by class; the activation grids copy their fills from the
  slide's own legend so a cell and the legend can never disagree.
* ``OPS_IMG:<key>`` pictures take the uploaded image, scaled to fit inside the
  template's frame and centred (never stretched); a slot with no image is
  removed rather than printing a placeholder.
* ``OPS_OPTIONAL:<key>`` slides that are switched off are deleted, with any
  slide-jump link to them removed so the file opens without repair.
"""
import copy
import datetime
import io
import logging
import re

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.image import Image

from app import opsum

log = logging.getLogger(__name__)

R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

TOKEN = re.compile(r"\{\{([a-z0-9_]+(?:\.\d+)*)\}\}")
WHOLE_TOKEN = re.compile(r"^\s*\{\{([a-z0-9_]+)\}\}\s*$")
INLINE = re.compile(r"\*\*(.+?)\*\*|\[([^\]\n]+)\]\((https?://[^)\s]+)\)")
SUB_INDENT = 342900        # EMU between bullet levels in the source deck

# Children of a:rPr that must come AFTER hlinkClick (schema order).
_RPR_AFTER_LINK = (qn("a:hlinkMouseOver"), qn("a:rtl"), qn("a:extLst"))
# Children of a:tcPr that come BEFORE the fill (schema order).
_TCPR_BEFORE_FILL = tuple(qn(f"a:{n}") for n in
                          ("lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D"))
_FILLS = tuple(qn(f"a:{n}") for n in
               ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"))


# --------------------------------------------------------------------------- #
# Inline markup
# --------------------------------------------------------------------------- #
def segments(text):
    """``"a **b** [c](https://x)"`` -> [("a ", False, None), ("b", True, None),
    (" ", False, None), ("c", False, "https://x")].

    Bold is None throughout when the text uses no ``**``, meaning "keep the
    template's weight" — so a bold count stays bold. Once the text does use
    ``**``, the markup decides every segment (a bold template cell holding
    ``**35,598** (18.2%)`` must print the percentage regular)."""
    out, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, None))
        if m.group(1) is not None:
            out.append((m.group(1), True, None))
        else:
            out.append((m.group(2), False, m.group(3)))
        pos = m.end()
    if pos < len(text) or not out:
        out.append((text[pos:], False, None))
    if not any(b for _, b, _ in out):
        out = [(t, None, u) for t, _, u in out]
    return out


def _make_run(rpr, text, bold, url, part):
    r = etree.Element(qn("a:r"))
    rpr = copy.deepcopy(rpr) if rpr is not None else etree.Element(qn("a:rPr"))
    for h in rpr.findall(qn("a:hlinkClick")):
        rpr.remove(h)
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if url:
        link = etree.Element(qn("a:hlinkClick"))
        link.set(R_ID, part.relate_to(url, RT.HYPERLINK, is_external=True))
        after = [c for c in rpr if c.tag in _RPR_AFTER_LINK]
        if after:
            after[0].addprevious(link)
        else:
            rpr.append(link)
    r.append(rpr)
    t = etree.SubElement(r, qn("a:t"))
    t.text = text
    return r


def _runs(p):
    return p.findall(qn("a:r"))


def _first_rpr(p):
    for r in _runs(p):
        rpr = r.find(qn("a:rPr"))
        if rpr is not None:
            return rpr
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        rpr = etree.Element(qn("a:rPr"))
        for k, v in end.attrib.items():
            rpr.set(k, v)
        for c in end:
            rpr.append(copy.deepcopy(c))
        return rpr
    return None


def _set_runs(p, text, rpr, part):
    for c in list(p):
        if c.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p.remove(c)
    end = p.find(qn("a:endParaRPr"))
    for seg_text, bold, url in segments(text):
        if not seg_text:
            continue
        r = _make_run(rpr, seg_text, bold, url, part)
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


# --------------------------------------------------------------------------- #
# Paragraph expansion (text fields)
# --------------------------------------------------------------------------- #
def _ppr(p):
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr"))
        p.insert(0, ppr)
    return ppr


def _no_bullet(ppr):
    for tag in ("a:buChar", "a:buAutoNum", "a:buBlip", "a:buNone"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    none = etree.Element(qn("a:buNone"))
    # buNone sits with the other bu* elements, before tabLst/defRPr/extLst.
    tail = [c for c in ppr if c.tag in (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))]
    if tail:
        tail[0].addprevious(none)
    else:
        ppr.append(none)
    ppr.set("marL", "0")
    ppr.set("indent", "0")


def _heading_proto(p):
    """A heading paragraph from the same cell/table/text box as ``p`` —
    unbulleted with a bold first run, like "Last 24 Hours" — so ``# Heading``
    lines match the slide's own headings (colour, size, spacing)."""
    tc = next(p.iterancestors(qn("a:tc")), None)
    tbl = next(p.iterancestors(qn("a:tbl")), None)
    # Own cell first, then the table's other body cells — row 0 is the
    # table's title bar ("PUBLIC HEALTH"), a different style.
    scopes = [tc] if tc is not None else [p.getparent()]
    if tbl is not None:
        scopes += tbl.findall(qn("a:tr"))[1:]
    for scope in scopes:
        for q in scope.iter(qn("a:p")):
            text = "".join(t.text or "" for t in q.iter(qn("a:t")))
            if q is p or "{{" in text or not text.strip():
                continue
            ppr = q.find(qn("a:pPr"))
            bulleted = ppr is not None and (ppr.find(qn("a:buChar")) is not None
                                            or ppr.find(qn("a:buAutoNum")) is not None)
            rpr = _first_rpr(q)
            if not bulleted and rpr is not None and rpr.get("b") == "1":
                return q
    return None


def _size_empty(p, rpr):
    """An empty paragraph is as tall as its endParaRPr size — the 18 pt
    default when it has none, which blows table rows and spacers up. Give it
    the size of the text it stands in for."""
    if rpr is None or not rpr.get("sz"):
        return
    end = p.find(qn("a:endParaRPr"))
    if end is None:
        end = etree.SubElement(p, qn("a:endParaRPr"))
    end.set("sz", rpr.get("sz"))


def expand(p, value, part):
    """Replace token paragraph ``p`` with one paragraph per line of ``value``."""
    lines = (value or "").split("\n")
    while len(lines) > 1 and not lines[-1].strip():
        lines.pop()
    rpr = _first_rpr(p)
    heading = None
    for line in lines:
        stripped = line.strip()
        sub = bool(stripped) and line[:1] in (" ", "\t")
        stripped = re.sub(r"^(?:[-•*]\s+)", "", stripped)
        if stripped.startswith("# "):
            if heading is None:
                found = _heading_proto(p)
                heading = found if found is not None else False
            if heading is not False:
                new = copy.deepcopy(heading)
                _set_runs(new, stripped[2:].strip(), _first_rpr(heading), part)
            else:
                new = copy.deepcopy(p)
                _no_bullet(_ppr(new))
                run_rpr = copy.deepcopy(rpr) if rpr is not None else None
                if run_rpr is not None and run_rpr.get("sz", "").isdigit():
                    run_rpr.set("sz", str(int(run_rpr.get("sz")) + 100))
                _set_runs(new, "**" + stripped[2:].strip() + "**", run_rpr, part)
            p.addprevious(new)
            continue
        new = copy.deepcopy(p)
        ppr = _ppr(new)
        if not stripped:
            _no_bullet(ppr)
            _set_runs(new, "", rpr, part)
            _size_empty(new, rpr)
        else:
            if sub:
                ppr.set("lvl", "1")
                ppr.set("marL", str(int(ppr.get("marL", "0") or 0) + SUB_INDENT))
            _set_runs(new, stripped, rpr, part)
        p.addprevious(new)
    p.getparent().remove(p)


# --------------------------------------------------------------------------- #
# Cells
# --------------------------------------------------------------------------- #
def set_cell_fill(tc, fill):
    """``fill`` is an RGB hex string, an ``a:*Fill`` element to copy, or None
    to fall back to the table style."""
    tcpr = tc.find(qn("a:tcPr"))
    if tcpr is None:
        tcpr = etree.SubElement(tc, qn("a:tcPr"))
    for el in list(tcpr):
        if el.tag in _FILLS:
            tcpr.remove(el)
    if fill is None:
        return
    if isinstance(fill, str):
        el = etree.Element(qn("a:solidFill"))
        etree.SubElement(el, qn("a:srgbClr")).set("val", fill)
    else:
        el = copy.deepcopy(fill)
    before = [c for c in tcpr if c.tag in _TCPR_BEFORE_FILL]
    if before:
        before[-1].addnext(el)
    else:
        tcpr.insert(0, el)


def _cell_text(tc):
    return "".join(t.text or "" for t in tc.iter(qn("a:t")))


def _cell_fill(tc):
    tcpr = tc.find(qn("a:tcPr"))
    if tcpr is None:
        return None
    for el in tcpr:
        if el.tag in _FILLS:
            return el
    return None


def _norm(label):
    return re.sub(r"\s+", " ", label or "").strip().lower()


def apply_grids(slide, fields):
    legend = None
    for shape in slide.shapes:
        if shape.name == "OPS_GRID_LEGEND" and shape.has_table:
            legend = {}
            for row in shape.table.rows:
                for cell in row.cells:
                    legend[_norm(cell.text)] = _cell_fill(cell._tc)
    for shape in slide.shapes:
        if not (shape.name.startswith("OPS_GRID:") and shape.has_table):
            continue
        key = shape.name.split(":", 1)[1]
        states = {_norm(k): v for k, v in (fields.get(key) or {}).items()}
        for row in shape.table.rows:
            for cell in row.cells:
                label = _norm(cell.text)
                if not label:
                    continue
                state = states.get(label, opsum.GRID_STATES[0])
                set_cell_fill(cell._tc, (legend or {}).get(_norm(state)))


def apply_cell_colours(part, fields):
    """AV workload cells take a colour and lose their token; flood Status
    cells are coloured by class. Must run before the text pass."""
    for tc in list(part._element.iter(qn("a:tc"))):
        m = TOKEN.fullmatch(_cell_text(tc).strip())
        if not m:
            continue
        key, *idx = m.group(1).split(".")
        f = opsum.FIELDS.get(key)
        if f is None or len(idx) != 2:
            continue
        i, j = int(idx[0]), int(idx[1])
        try:
            value = fields[key][i][j]
        except (KeyError, IndexError, TypeError):
            continue
        if f.kind == "colors":
            colour = opsum.WORKLOAD_COLOURS.get(value)
            if colour:
                set_cell_fill(tc, colour)
            for t in tc.iter(qn("a:t")):
                t.text = ""
        elif key == "flood_gauges" and i == 2:
            colour = opsum.FLOOD_STATUS_FILLS.get((value or "").strip().upper())
            if colour:
                set_cell_fill(tc, colour)
                # The template's per-cell text colours came from whatever class
                # sat there in the source deck; black reads on all three fills.
                for rpr in tc.iter(qn("a:rPr")):
                    for el in rpr.findall(qn("a:solidFill")):
                        rpr.remove(el)
                    black = etree.Element(qn("a:solidFill"))
                    etree.SubElement(black, qn("a:srgbClr")).set("val", "000000")
                    ln = rpr.find(qn("a:ln"))
                    if ln is not None:
                        ln.addnext(black)
                    else:
                        rpr.insert(0, black)


# --------------------------------------------------------------------------- #
# Text pass
# --------------------------------------------------------------------------- #
def inline_values(d, fields):
    vals = dict(opsum.derived_tokens(d))
    for key, f in opsum.FIELDS.items():
        v = fields.get(key)
        if f.kind == "line":
            vals[key] = v or ""
        elif f.kind in ("table", "colors"):
            for i, row in enumerate(v or []):
                for j, cell in enumerate(row):
                    vals[f"{key}.{i}.{j}"] = cell or ""
        elif f.kind == "text":
            # A text field used inline (never in the shipped template) prints
            # its first line rather than leaking a token.
            vals[key] = (v or "").split("\n")[0]
    return vals


def fill_text(part, fields, values):
    for p in list(part._element.iter(qn("a:p"))):
        text = "".join(t.text or "" for t in p.iter(qn("a:t")))
        if "{{" not in text:
            continue
        m = WHOLE_TOKEN.match(text)
        if m and opsum.FIELDS.get(m.group(1)) and opsum.FIELDS[m.group(1)].kind == "text":
            expand(p, fields.get(m.group(1)) or "", part)
            continue
        for r in _runs(p):
            t = r.find(qn("a:t"))
            if t is None or not t.text or "{{" not in t.text:
                continue
            new = TOKEN.sub(lambda mm: _lookup(values, mm.group(1)), t.text)
            segs = segments(new)
            if len(segs) == 1 and segs[0][1] is None and not segs[0][2]:
                t.text = new
                if not new:
                    _size_empty(p, r.find(qn("a:rPr")))
                continue
            rpr = r.find(qn("a:rPr"))
            for seg_text, bold, url in segs:
                if seg_text:
                    r.addprevious(_make_run(rpr, seg_text, bold, url, part))
            p.remove(r)


def _lookup(values, token):
    if token not in values:
        log.warning("Operational Summary template token {{%s}} has no value", token)
        return ""
    return values[token]


# --------------------------------------------------------------------------- #
# Images and slides
# --------------------------------------------------------------------------- #
def apply_images(slide, fields):
    def walk(shapes):
        for s in shapes:
            if s.shape_type == 6:
                yield from walk(s.shapes)
            elif s.name.startswith("OPS_IMG:"):
                yield s

    for pic in list(walk(slide.shapes)):
        key = pic.name.split(":", 1)[1]
        path = opsum.image_path(fields.get(key))
        blip = pic._element.blipFill.find(qn("a:blip"))
        old = blip.get(R_EMBED)
        if not path:
            pic._element.getparent().remove(pic._element)
            slide.part.drop_rel(old)
            continue
        with open(path, "rb") as fh:
            blob = fh.read()
        _, rid = slide.part.get_or_add_image_part(io.BytesIO(blob))
        blip.set(R_EMBED, rid)
        slide.part.drop_rel(old)
        src = pic._element.blipFill.find(qn("a:srcRect"))
        if src is not None:
            pic._element.blipFill.remove(src)
        w, h = Image.from_blob(blob).size
        if w and h:
            x, y, cx, cy = pic.left, pic.top, pic.width, pic.height
            scale = min(cx / w, cy / h)
            nw, nh = int(w * scale), int(h * scale)
            pic.left, pic.top = x + (cx - nw) // 2, y + (cy - nh) // 2
            pic.width, pic.height = nw, nh
        pic._element.nvPicPr.cNvPr.set("descr", (fields.get(key) or {}).get("name") or key)


def drop_slide(prs, slide):
    gone = slide.part
    for other in prs.slides:
        if other.part is gone:
            continue
        for rid, rel in list(other.part.rels.items()):
            if rel.is_external or rel.target_part is not gone:
                continue
            for el in list(other.part._element.iter(qn("a:hlinkClick"))):
                if el.get(R_ID) == rid:
                    el.getparent().remove(el)
            other.part.drop_rel(rid)
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        if prs.part.related_part(sld.rId) is gone:
            rid = sld.rId
            lst.remove(sld)
            prs.part.drop_rel(rid)


def render(d, data, template=None):
    """Build the deck for summary date ``d`` from ``data``; returns bytes."""
    d = opsum.parse_date(d)
    data = opsum.normalise(data, d)
    fields = data["fields"]
    prs = Presentation(template or opsum.TEMPLATE_PATH)

    for slide in list(prs.slides):
        name = slide._element.cSld.get("name") or ""
        if name.startswith("OPS_OPTIONAL:") and not data["optional"].get(name.split(":", 1)[1]):
            drop_slide(prs, slide)

    values = inline_values(d, fields)
    parts = [s.part for s in prs.slides] + [l.part for l in prs.slide_layouts]
    for slide in prs.slides:
        apply_grids(slide, fields)
        apply_images(slide, fields)
    for part in parts:
        apply_cell_colours(part, fields)
        fill_text(part, fields, values)

    cp = prs.core_properties
    cp.title = f"State Operational Summary — {d:%d %B %Y}"
    cp.modified = datetime.datetime.now()
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
